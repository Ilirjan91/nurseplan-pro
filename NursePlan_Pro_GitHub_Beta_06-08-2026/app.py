import calendar
import json
import re
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from time import monotonic, perf_counter
from urllib.parse import urlencode
from uuid import uuid4
from xml.sax.saxutils import escape

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
import extra_streamlit_components as stx

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from database import (
    SupabaseConfigurationError,
    add_feedback,
    add_mitarbeitende,
    add_vorgabe,
    create_tables,
    clear_local_auth_state,
    current_session_refresh_token,
    current_user_email,
    delete_dienstplan,
    delete_mitarbeitende,
    delete_vorgabe,
    get_dienstplaene,
    get_einstellungen,
    get_feedback,
    get_mitarbeitende,
    get_supabase_public_config,
    get_supabase_client,
    get_vorgaben,
    get_vorlagen,
    is_authenticated,
    load_dienstplan,
    load_vorlage,
    restore_user_session,
    request_password_reset,
    save_dienstplan,
    save_einstellungen,
    save_vorlage,
    delete_vorlage,
    sign_in_user,
    sign_out_user,
    sign_up_user,
    update_mitarbeitende,
)

st.set_page_config(page_title="NursePlan Pro", page_icon="🏥", layout="wide")

APP_DIRECTORY = Path(__file__).resolve().parent
PAINT_GRID_CANDIDATES = [
    APP_DIRECTORY / "paint_grid" / "dist",
    APP_DIRECTORY.parent / "paint_grid" / "dist",
]
PAINT_GRID_BUILD = next(
    (candidate for candidate in PAINT_GRID_CANDIDATES if candidate.is_dir()),
    None,
)
paint_grid_component = (
    components.declare_component(
        "nurseplan_paint_grid",
        path=str(PAINT_GRID_BUILD),
    )
    if PAINT_GRID_BUILD is not None
    else None
)

AUTH_COOKIE_NAME = "nurseplan_pro_session_v1"
AUTH_COOKIE_DAYS = 30

QUALIFIKATIONEN = [
    "Pflegefachkraft",
    "Praxisanleitung",
    "Stroke Nurse",
    "Schüler/in",
    "Pflegehelfer/in",
    "Stationshelfer/in",
    "Stationsleitung",
    "Stationssekretär/in",
]

FESTE_DIENSTARTEN = [
    "Alle Dienste",
    "Nur Frühdienst",
    "Nur Spätdienst",
    "Nur Nachtdienst",
]

FESTE_DIENSTART_CODES = {
    "Alle Dienste": {"F", "S", "N"},
    "Nur Frühdienst": {"F"},
    "Nur Spätdienst": {"S"},
    "Nur Nachtdienst": {"N"},
}

DIENST_CODES = ["", "F", "S", "N", "U", "W", "FB", "K", "Frei"]
WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
MONATSNAMEN = {
    1: "Januar", 2: "Februar", 3: "März", 4: "April",
    5: "Mai", 6: "Juni", 7: "Juli", 8: "August",
    9: "September", 10: "Oktober", 11: "November", 12: "Dezember",
}

DEFAULT_SETTINGS = {
    "vollzeit_monatssoll": 167.0,
    "f_start": "06:00",
    "f_ende": "14:12",
    "f_stunden": 7.7,
    "min_f": 2,
    "s_start": "13:48",
    "s_ende": "22:00",
    "s_stunden": 7.7,
    "min_s": 2,
    "n_start": "21:48",
    "n_ende": "06:00",
    "n_stunden": 7.7,
    "min_n": 1,
    "min_fach_f": 1,
    "min_fach_s": 1,
    "min_fach_n": 1,
    "ruhezeit": 11.0,
    "max_arbeitstage": 6,
    "max_nachtdienste": 4,
    "max_wochenenden": 2,
    "keine_folgewochenenden": True,
    "max_plusstunden": 0.0,
    "regel_ruhezeit_aktiv": True,
    "regel_max_arbeitstage_aktiv": True,
    "regel_max_nachtdienste_aktiv": True,
    "regel_max_wochenenden_aktiv": True,
    "regel_sollstunden_aktiv": True,
    "regel_fachkraft_aktiv": True,
    "regel_nacht_zuerst_aktiv": True,
    "regel_nachtblock_aktiv": True,
    "regel_wochenende_zuerst_aktiv": True,
    "regel_wochenende_gleich_aktiv": True,
    "eigene_planungsregeln": [],
}

ACCOUNT_CACHE_TTL_SECONDS = 300.0
ACCOUNT_CACHE_STATE_KEY = "_np_account_data_cache"
PERFORMANCE_STATE_KEY = "_np_performance_metrics"


def record_performance(label: str, seconds: float, cached: bool = False) -> None:
    """Speichert nur technische Laufzeiten in der aktuellen Nutzersitzung."""
    metrics = st.session_state.setdefault(PERFORMANCE_STATE_KEY, {})
    metrics[str(label)] = {
        "milliseconds": round(max(0.0, float(seconds)) * 1000.0, 1),
        "cached": bool(cached),
        "measured_at": datetime.now().strftime("%H:%M:%S"),
    }


def measured_call(label: str, function, *args, **kwargs):
    started = perf_counter()
    try:
        return function(*args, **kwargs)
    finally:
        record_performance(label, perf_counter() - started)


def _account_cache() -> dict:
    """Liefert einen Cache, der ausschließlich zur aktuellen Anmeldung gehört."""
    owner_id = str(st.session_state.get("auth_user_id", ""))
    cache = st.session_state.get(ACCOUNT_CACHE_STATE_KEY)
    if not isinstance(cache, dict) or cache.get("owner_id") != owner_id:
        cache = {"owner_id": owner_id, "entries": {}}
        st.session_state[ACCOUNT_CACHE_STATE_KEY] = cache
    return cache


def invalidate_account_cache(*names: str) -> None:
    """Entfernt geänderte Datensätze sofort aus dem Sitzungs-Cache."""
    cache = st.session_state.get(ACCOUNT_CACHE_STATE_KEY)
    if not isinstance(cache, dict):
        return
    if not names:
        st.session_state.pop(ACCOUNT_CACHE_STATE_KEY, None)
        return
    entries = cache.setdefault("entries", {})
    for name in names:
        prefix = f"{name}/"
        for entry_name in list(entries):
            if entry_name == str(name) or entry_name.startswith(prefix):
                entries.pop(entry_name, None)


def load_account_dataset(name: str, loader):
    """Lädt Kontodaten höchstens alle fünf Minuten neu und niemals kontoübergreifend."""
    cache = _account_cache()
    entries = cache.setdefault("entries", {})
    now = monotonic()
    cached_entry = entries.get(name)
    if (
        isinstance(cached_entry, dict)
        and now - float(cached_entry.get("loaded_at", 0.0)) < ACCOUNT_CACHE_TTL_SECONDS
    ):
        record_performance(f"Datenbank · {name}", 0.0, cached=True)
        return deepcopy(cached_entry.get("value"))

    started = perf_counter()
    value = loader()
    record_performance(f"Datenbank · {name}", perf_counter() - started)
    entries[name] = {"loaded_at": now, "value": deepcopy(value)}
    return value


def load_account_overview_data():
    dataset_names = ("Mitarbeitende", "Abwesenheiten", "Dienstpläne", "Vorlagen")
    cache_entries = _account_cache().setdefault("entries", {})
    checked_at = monotonic()
    all_cached = all(
        isinstance(cache_entries.get(name), dict)
        and checked_at - float(cache_entries[name].get("loaded_at", 0.0))
        < ACCOUNT_CACHE_TTL_SECONDS
        for name in dataset_names
    )
    started = perf_counter()
    result = (
        load_account_dataset("Mitarbeitende", get_mitarbeitende),
        load_account_dataset("Abwesenheiten", get_vorgaben),
        load_account_dataset("Dienstpläne", get_dienstplaene),
        load_account_dataset("Vorlagen", get_vorlagen),
    )
    record_performance("Kontodaten gesamt", perf_counter() - started, cached=all_cached)
    return result

st.markdown(
    """
    <style>
    :root {
        --np-navy: #16324f;
        --np-blue: #2563eb;
        --np-blue-soft: #eaf2ff;
        --np-green: #0f9f6e;
        --np-orange: #d97706;
        --np-red: #dc2626;
        --np-text: #172033;
        --np-muted: #667085;
        --np-line: #e5e9f0;
        --np-bg: #f4f7fb;
        --np-card: #ffffff;
    }

    .stApp {
        background: var(--np-bg);
        color: var(--np-text);
        color-scheme: light;
    }
    header[data-testid="stHeader"] {
        background: rgba(244, 247, 251, .98) !important;
        border-bottom: 1px solid var(--np-line);
    }
    [data-testid="stMainBlockContainer"] {
        max-width: 1600px;
        padding: 1.25rem 1.6rem 3rem 1.6rem;
    }

    section[data-testid="stSidebar"] {
        background: linear-gradient(180deg, #102a43 0%, #173d5f 100%);
        border-right: 0;
    }
    section[data-testid="stSidebar"] * { color: #f8fafc !important; }
    section[data-testid="stSidebar"] [data-testid="stRadio"] label {
        padding: .6rem .75rem;
        border-radius: 10px;
        margin-bottom: .15rem;
    }
    section[data-testid="stSidebar"] [data-testid="stRadio"] label:hover {
        background: rgba(255,255,255,.09);
    }

    .np-brand {
        padding: .25rem 0 1rem 0;
        border-bottom: 1px solid rgba(255,255,255,.16);
        margin-bottom: .75rem;
    }
    .np-brand-name { font-size: 1.25rem; font-weight: 800; letter-spacing: -.02em; }
    .np-brand-subtitle { font-size: .82rem; opacity: .76; margin-top: .2rem; }

    .np-title {
        font-size: 2.15rem;
        line-height: 1.15;
        font-weight: 800;
        letter-spacing: -.035em;
        margin: 0;
        color: var(--np-text);
    }
    .np-subtitle {
        color: var(--np-muted);
        font-size: 1rem;
        margin: .35rem 0 1.3rem 0;
        max-width: 900px;
    }
    .np-section-title {
        font-size: 1.12rem;
        font-weight: 750;
        margin: .2rem 0 .2rem 0;
    }
    .np-help { color: var(--np-muted); font-size: .9rem; margin-bottom: .8rem; }

    .np-card {
        background: var(--np-card);
        border: 1px solid var(--np-line);
        border-radius: 16px;
        padding: 1.05rem 1.15rem;
        margin-bottom: 1rem;
        box-shadow: 0 4px 16px rgba(15, 23, 42, .035);
    }
    .np-kpi {
        background: var(--np-card);
        border: 1px solid var(--np-line);
        border-radius: 15px;
        padding: 1rem 1.05rem;
        min-height: 112px;
        box-shadow: 0 4px 16px rgba(15, 23, 42, .035);
    }
    .np-kpi-label { color: var(--np-muted); font-size: .82rem; font-weight: 650; }
    .np-kpi-value { color: var(--np-text); font-size: 1.8rem; font-weight: 800; margin-top: .15rem; }
    .np-kpi-note { color: var(--np-muted); font-size: .78rem; margin-top: .12rem; }

    .np-percent {
        background: var(--np-blue-soft);
        border: 1px solid #cfddff;
        border-radius: 12px;
        padding: .7rem .9rem;
        margin: .35rem 0 .8rem 0;
    }
    .np-percent-label { color: var(--np-muted); font-size: .8rem; }
    .np-percent-value { color: #1d4ed8; font-size: 1.35rem; font-weight: 800; }

    .np-table-wrap {
        overflow-x: auto;
        border: 1px solid var(--np-line);
        border-radius: 14px;
        background: white;
    }
    .np-table { width: 100%; border-collapse: collapse; min-width: 780px; }
    .np-table th {
        text-align: left;
        background: #f8fafc;
        color: #475467;
        font-size: .78rem;
        text-transform: uppercase;
        letter-spacing: .035em;
        padding: .72rem .8rem;
        border-bottom: 1px solid var(--np-line);
    }
    .np-table td {
        padding: .72rem .8rem;
        border-bottom: 1px solid #eef1f5;
        font-size: .9rem;
        vertical-align: middle;
    }
    .np-table tr:last-child td { border-bottom: 0; }
    .np-table tr:hover td { background: #fbfdff; }
    .np-name { font-weight: 750; color: var(--np-text); }
    .np-badge {
        display: inline-block;
        border-radius: 999px;
        padding: .2rem .55rem;
        font-size: .75rem;
        font-weight: 700;
        white-space: nowrap;
    }
    .np-badge-blue { background: #eaf2ff; color: #1d4ed8; }
    .np-badge-green { background: #e8f8f2; color: #087a54; }
    .np-badge-gray { background: #f1f3f6; color: #475467; }
    .np-badge-orange { background: #fff4df; color: #a85b00; }
    .np-badge-red { background: #feecec; color: #b42318; }

    .np-template-card {
        background: white;
        border: 1px solid var(--np-line);
        border-radius: 15px;
        padding: 1rem;
        min-height: 195px;
        margin-bottom: .45rem;
    }
    .np-template-title { font-weight: 800; font-size: 1.02rem; margin-bottom: .35rem; }
    .np-template-text { color: var(--np-muted); font-size: .86rem; line-height: 1.45; }

    div[data-testid="stDataFrame"], div[data-testid="stDataEditor"] {
        background: white;
        border: 1px solid var(--np-line);
        border-radius: 14px;
        overflow: hidden;
    }
    div[data-testid="stForm"] {
        background: white;
        border: 1px solid var(--np-line);
        border-radius: 15px;
        padding: 1.15rem;
    }
    [data-testid="stMain"] [data-testid="stWidgetLabel"] p,
    [data-testid="stMain"] div[data-testid="stTextInput"] label p,
    [data-testid="stMain"] div[data-testid="stNumberInput"] label p,
    [data-testid="stMain"] div[data-testid="stSelectbox"] label p,
    [data-testid="stMain"] div[data-testid="stDateInput"] label p,
    [data-testid="stMain"] div[data-testid="stTextArea"] label p {
        color: #344054 !important;
        font-size: .88rem !important;
        font-weight: 700 !important;
    }
    [data-testid="stMain"] div[data-baseweb="input"],
    [data-testid="stMain"] div[data-baseweb="base-input"],
    [data-testid="stMain"] div[data-baseweb="input"] > div {
        background: #ffffff !important;
        border-color: #cfd6e1 !important;
        color: #172033 !important;
    }
    [data-testid="stMain"] div[data-baseweb="input"] {
        border: 1px solid #cfd6e1 !important;
        border-radius: 10px !important;
        box-shadow: 0 1px 2px rgba(16, 24, 40, .04);
    }
    [data-testid="stMain"] div[data-baseweb="input"]:focus-within {
        border-color: var(--np-blue) !important;
        box-shadow: 0 0 0 3px rgba(37, 99, 235, .12) !important;
    }
    [data-testid="stMain"] div[data-baseweb="input"] input {
        color: #172033 !important;
        -webkit-text-fill-color: #172033 !important;
        caret-color: var(--np-blue) !important;
    }
    [data-testid="stMain"] div[data-baseweb="input"] input::placeholder {
        color: #98a2b3 !important;
        -webkit-text-fill-color: #98a2b3 !important;
    }
    [data-testid="stMain"] div[data-baseweb="input"] svg {
        color: #667085 !important;
        fill: #667085 !important;
    }
    [data-testid="stMain"] div[data-baseweb="input"] button {
        background: transparent !important;
        color: #667085 !important;
    }
    div[data-testid="stExpander"] {
        border: 1px solid var(--np-line);
        border-radius: 12px;
        background: white;
    }
    .stButton button, .stDownloadButton button, .stFormSubmitButton button {
        border-radius: 10px;
        font-weight: 700;
        min-height: 2.55rem;
    }
    button[kind="primary"] {
        background: var(--np-blue) !important;
        border-color: var(--np-blue) !important;
        color: #ffffff !important;
        box-shadow: 0 3px 8px rgba(37, 99, 235, .18);
    }
    button[kind="primary"] p { color: #ffffff !important; }
    button[kind="primary"]:hover {
        background: #1d4ed8 !important;
        border-color: #1d4ed8 !important;
    }
    .stTabs [data-baseweb="tab-list"] { gap: .35rem; }
    .stTabs [data-baseweb="tab"] {
        border-radius: 9px 9px 0 0;
        padding: .55rem .9rem;
        font-weight: 700;
        color: #667085 !important;
    }
    .stTabs [data-baseweb="tab"][aria-selected="true"] {
        color: #1d4ed8 !important;
    }
    .stTabs [data-baseweb="tab"] p { color: inherit !important; }
    .stTabs [data-baseweb="tab-highlight"] {
        background-color: var(--np-blue) !important;
    }
    hr { border-color: var(--np-line); }
    </style>
    """,
    unsafe_allow_html=True,
)


auth_cookie_manager = stx.CookieManager(key="nurseplan_auth_cookie_manager")


def auth_cookie_is_secure() -> bool:
    """Secure-Cookies auf HTTPS; localhost bleibt für die Entwicklung nutzbar."""
    try:
        return str(st.context.url).casefold().startswith("https://")
    except Exception:
        return False


def saved_auth_cookie() -> str:
    value = auth_cookie_manager.get(AUTH_COOKIE_NAME)
    if not value:
        return ""
    # Der Browser-Cookie-Manager wandelt JSON-Cookies automatisch in ein
    # Python-Dictionary um. Für die weitere Verarbeitung wieder als gültiges
    # JSON serialisieren; str(dict) wäre kein gültiges JSON und würde sonst
    # fälschlich als Refresh-Token behandelt.
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).strip()


def adopt_google_user_session(access_token: str, refresh_token: str) -> str:
    """Übernimmt eine neue Google-Sitzung, ohne den Refresh-Token zu rotieren."""
    response = get_supabase_client().auth.set_session(access_token, refresh_token)
    user = getattr(response, "user", None)
    session = getattr(response, "session", None)
    if user is None or session is None:
        raise RuntimeError("Die Google-Sitzung konnte nicht übernommen werden.")

    metadata = getattr(user, "user_metadata", None) or {}
    st.session_state.auth_user_id = str(user.id)
    st.session_state.auth_email = str(getattr(user, "email", "") or "")
    st.session_state.auth_name = str(metadata.get("full_name", "") or "")

    active_refresh_token = str(getattr(session, "refresh_token", "") or "").strip()
    if not active_refresh_token:
        raise RuntimeError("Die Google-Sitzung enthält keinen Refresh-Token.")
    return active_refresh_token


def saved_auth_tokens() -> tuple[str, str]:
    """Liest entweder eine neue OAuth-Sitzung oder einen normalen Refresh-Token."""
    raw_value = saved_auth_cookie()
    if not raw_value:
        return "", ""
    # Neue Google-Sitzungen werden bewusst als einfacher Text gespeichert.
    # So kann der Browser-Cookie-Manager den Inhalt nicht automatisch in ein
    # Objekt umwandeln. Access- und Refresh-Token enthalten kein Pipe-Zeichen.
    if raw_value.startswith("oauth_v1|"):
        parts = raw_value.split("|", 2)
        if len(parts) == 3 and parts[1] and parts[2]:
            return parts[1].strip(), parts[2].strip()
        return "", ""
    try:
        oauth_session = json.loads(raw_value)
    except (TypeError, ValueError, json.JSONDecodeError):
        return "", raw_value
    if not isinstance(oauth_session, dict):
        return "", raw_value
    return (
        str(oauth_session.get("access_token", "") or "").strip(),
        str(oauth_session.get("refresh_token", "") or "").strip(),
    )


def save_auth_cookie(refresh_token: str) -> None:
    token = str(refresh_token or "").strip()
    if not token or saved_auth_cookie() == token:
        return
    auth_cookie_manager.set(
        AUTH_COOKIE_NAME,
        token,
        key="save_nurseplan_auth_cookie",
        path="/",
        expires_at=datetime.now() + timedelta(days=AUTH_COOKIE_DAYS),
        max_age=AUTH_COOKIE_DAYS * 24 * 60 * 60,
        secure=auth_cookie_is_secure(),
        same_site="strict",
    )


def delete_auth_cookie() -> None:
    if not saved_auth_cookie():
        return
    auth_cookie_manager.delete(
        AUTH_COOKIE_NAME,
        key="delete_nurseplan_auth_cookie",
    )


def auth_error_text(error: Exception, action: str) -> str:
    message = str(error).casefold()
    if "invalid login credentials" in message:
        return "E-Mail oder Passwort ist nicht richtig."
    if "email not confirmed" in message:
        return "Bitte bestätige zuerst deine E-Mail-Adresse."
    if "already registered" in message:
        return "Für diese E-Mail-Adresse besteht bereits ein Konto."
    if "rate limit" in message or "too many" in message:
        return "Zu viele Versuche. Bitte warte kurz und versuche es erneut."
    if "password" in message and ("least" in message or "short" in message):
        return "Das Passwort ist zu kurz. Verwende mindestens 8 Zeichen."
    if action == "login":
        return "Die Anmeldung ist gerade nicht möglich. Bitte versuche es erneut."
    if action == "password_reset":
        return "Die E-Mail konnte gerade nicht gesendet werden. Bitte versuche es später erneut."
    return "Die Registrierung ist gerade nicht möglich. Bitte versuche es erneut."


def google_oauth_session_error_text(error: Exception) -> str:
    """Zeigt einen hilfreichen Hinweis, ohne Token oder andere Secrets auszugeben."""
    message = str(error).casefold()
    if "refresh token" in message:
        return (
            "Supabase hat die Google-Sitzung nicht angenommen "
            "(Refresh-Token ungültig oder nicht gefunden)."
        )
    if "jwt" in message or "access token" in message or "user" in message:
        return "Supabase konnte das Google-Konto nicht bestätigen (Zugriffstoken ungültig)."
    return (
        "Die Google-Sitzung konnte serverseitig nicht übernommen werden "
        f"({type(error).__name__})."
    )


def password_reset_redirect_url() -> str:
    app_url = configured_app_url()
    return f"{app_url.rstrip('/')}/?password_recovery=1"


def configured_app_url() -> str:
    """Liefert die öffentliche App-Adresse ohne Pfad oder abschließenden Slash."""
    try:
        configured_url = str(st.secrets.get("APP_URL", "") or "").strip()
    except Exception:
        configured_url = ""
    return (configured_url or "http://localhost:8501").rstrip("/")


def google_oauth_redirect_url() -> str:
    return f"{configured_app_url()}/?oauth_callback=google"


def google_oauth_url() -> str:
    """Erstellt einen Browser-OAuth-Link im Supabase-Implicit-Flow.

    Der Python-Client erzeugt sonst einen PKCE-Code. Da Streamlit den Link in
    einem neuen Browser-Tab öffnet, fehlt dort der ursprüngliche PKCE-Verifier.
    Ohne Code-Challenge liefert Supabase die neue Sitzung sicher im URL-Fragment,
    das anschließend ausschließlich im Browser verarbeitet wird.
    """
    supabase_url, _ = get_supabase_public_config()
    query = urlencode(
        {
            "provider": "google",
            "redirect_to": google_oauth_redirect_url(),
        }
    )
    return f"{supabase_url}/auth/v1/authorize?{query}"


def show_google_oauth_callback(show_status: bool = True) -> None:
    """Übernimmt die neue Google-Sitzung kurzzeitig in das App-Cookie."""
    secure_cookie = "true" if auth_cookie_is_secure() else "false"
    callback_html = r"""
<!doctype html>
<html lang="de">
<head><meta charset="utf-8"></head>
<body>
<script>
(() => {
  // Streamlit Cloud zeigt die App in einem zusätzlichen äußeren Fenster an.
  // Supabase schreibt die Sitzung in dessen URL, nicht in die URL des
  // eingebetteten HTML-Hilfsfensters.
  let appWindow = window.parent;
  try {
    if (window.top && window.top.location.href) appWindow = window.top;
  } catch (_) {
    // Lokal gibt es kein zusätzliches Streamlit-Cloud-Fenster.
  }

  const fragment = new URLSearchParams(appWindow.location.hash.slice(1));
  const accessToken = fragment.get("access_token") || "";
  const refreshToken = fragment.get("refresh_token") || "";
  const secureCookie = __SECURE_COOKIE__;
  const callbackRequested = new URLSearchParams(appWindow.location.search)
    .get("oauth_callback") === "google";

  if (!accessToken || !refreshToken) {
    if (callbackRequested) {
      appWindow.history.replaceState({}, "", `${appWindow.location.pathname}?oauth_error=1`);
      appWindow.location.reload();
    }
    return;
  }

  const oauthSession = `oauth_v1|${accessToken}|${refreshToken}`;
  let cookie = "__COOKIE_NAME__=" + encodeURIComponent(oauthSession)
    + "; Path=/; Max-Age=120; SameSite=Lax";
  if (secureCookie) cookie += "; Secure";
  appWindow.document.cookie = cookie;
  appWindow.history.replaceState({}, "", `${appWindow.location.pathname}?oauth_complete=1`);
  appWindow.location.reload();
})();
</script>
</body>
</html>
"""
    callback_html = callback_html.replace("__SECURE_COOKIE__", secure_cookie)
    callback_html = callback_html.replace("__COOKIE_NAME__", AUTH_COOKIE_NAME)

    if not show_status:
        components.html(callback_html, height=1, scrolling=False)
        return

    left, middle, right = st.columns([1, 1.25, 1])
    with middle:
        st.markdown(
            """
            <div class="np-card" style="text-align:center;margin-top:3rem;">
                <div style="font-size:2.2rem;">🏥</div>
                <div style="font-size:1.55rem;font-weight:800;">Google-Anmeldung wird abgeschlossen</div>
                <div style="color:#667085;margin-top:.35rem;">Einen Moment bitte …</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        components.html(callback_html, height=1, scrolling=False)


def show_password_recovery_page() -> None:
    """Zeigt das neue Passwortfeld; das Passwort geht direkt an Supabase Auth."""
    supabase_url, publishable_key = get_supabase_public_config()
    recovery_html = r"""
<!doctype html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
  * { box-sizing: border-box; }
  body {
    margin: 0; padding: 8px; background: transparent; color: #172033;
    font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
  }
  .card {
    background: #fff; border: 1px solid #e5e9f0; border-radius: 16px;
    padding: 24px; box-shadow: 0 12px 32px rgba(16,42,67,.08);
  }
  h2 { margin: 0 0 6px; font-size: 22px; }
  .subtitle { margin: 0 0 20px; color: #667085; font-size: 14px; line-height: 1.45; }
  label { display: block; margin: 14px 0 6px; font-size: 14px; font-weight: 700; }
  input {
    width: 100%; border: 1px solid #cfd6e2; border-radius: 10px;
    padding: 12px 13px; font-size: 15px; background: #fff; color: #172033;
  }
  input:focus { outline: 3px solid #dbeafe; border-color: #2563eb; }
  button {
    width: 100%; margin-top: 18px; border: 0; border-radius: 10px;
    padding: 12px 16px; background: #2563eb; color: white;
    font-size: 15px; font-weight: 750; cursor: pointer;
  }
  button:disabled { opacity: .65; cursor: wait; }
  .message { display: none; margin-top: 14px; border-radius: 10px; padding: 11px 12px; font-size: 14px; }
  .error { display: block; background: #fef2f2; color: #b42318; }
  .success { display: block; background: #ecfdf3; color: #067647; }
</style>
</head>
<body>
  <div class="card">
    <h2>Neues Passwort festlegen</h2>
    <p class="subtitle">Wähle ein sicheres Passwort mit mindestens 8 Zeichen.</p>
    <form id="passwordForm">
      <label for="password">Neues Passwort</label>
      <input id="password" type="password" minlength="8" autocomplete="new-password" required>
      <label for="passwordRepeat">Passwort wiederholen</label>
      <input id="passwordRepeat" type="password" minlength="8" autocomplete="new-password" required>
      <button id="submitButton" type="submit">Passwort speichern</button>
    </form>
    <div id="message" class="message" role="status"></div>
  </div>
<script>
(() => {
  const supabaseUrl = __SUPABASE_URL__;
  const publishableKey = __PUBLISHABLE_KEY__;
  const form = document.getElementById("passwordForm");
  const message = document.getElementById("message");
  const submitButton = document.getElementById("submitButton");

  let accessToken = "";
  try {
    const fragment = new URLSearchParams(window.parent.location.hash.slice(1));
    accessToken = fragment.get("access_token") || "";
    if (fragment.get("type") !== "recovery") accessToken = "";
  } catch (_) {
    accessToken = "";
  }

  if (!accessToken) {
    form.style.display = "none";
    message.textContent = "Der Link ist ungültig oder abgelaufen. Bitte fordere einen neuen Link an.";
    message.className = "message error";
    return;
  }

  form.addEventListener("submit", async (event) => {
    event.preventDefault();
    message.className = "message";
    const password = document.getElementById("password").value;
    const repeated = document.getElementById("passwordRepeat").value;
    if (password.length < 8) {
      message.textContent = "Das Passwort muss mindestens 8 Zeichen haben.";
      message.className = "message error";
      return;
    }
    if (password !== repeated) {
      message.textContent = "Die beiden Passwörter stimmen nicht überein.";
      message.className = "message error";
      return;
    }

    submitButton.disabled = true;
    submitButton.textContent = "Wird gespeichert …";
    try {
      const response = await fetch(`${supabaseUrl}/auth/v1/user`, {
        method: "PUT",
        headers: {
          "apikey": publishableKey,
          "Authorization": `Bearer ${accessToken}`,
          "Content-Type": "application/json"
        },
        body: JSON.stringify({ password })
      });
      const data = await response.json().catch(() => ({}));
      if (!response.ok) throw new Error(data.msg || data.message || "Passwort konnte nicht geändert werden.");
      form.style.display = "none";
      message.textContent = "Dein Passwort wurde geändert. Klicke jetzt unten auf „Zurück zur Anmeldung“.";
      message.className = "message success";
    } catch (_) {
      message.textContent = "Das Passwort konnte nicht geändert werden. Bitte fordere einen neuen Link an.";
      message.className = "message error";
      submitButton.disabled = false;
      submitButton.textContent = "Passwort speichern";
    }
  });
})();
</script>
</body>
</html>
"""
    recovery_html = recovery_html.replace("__SUPABASE_URL__", json.dumps(supabase_url))
    recovery_html = recovery_html.replace("__PUBLISHABLE_KEY__", json.dumps(publishable_key))

    left, middle, right = st.columns([1, 1.25, 1])
    with middle:
        st.markdown(
            """
            <div class="np-card" style="text-align:center;margin-top:3rem;">
                <div style="font-size:2.2rem;">🏥</div>
                <div style="font-size:1.65rem;font-weight:800;">NursePlan Pro</div>
                <div style="color:#667085;margin-top:.25rem;">Sicheres Passwort festlegen</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        components.html(recovery_html, height=390, scrolling=False)
        if st.button(
            "← Zurück zur Anmeldung",
            type="secondary",
            width="stretch",
            key="password_recovery_back_to_login",
        ):
            delete_auth_cookie()
            clear_local_auth_state()
            st.session_state.auth_logout_requested = True
            st.query_params.clear()
            st.rerun()


def legal_config_value(name: str, default: str = "") -> str:
    try:
        return str(st.secrets.get(name, default) or default).strip()
    except Exception:
        return str(default).strip()


def legal_contact_details() -> dict[str, str]:
    return {
        "name": legal_config_value("LEGAL_OPERATOR_NAME", "Ilirjan Balisha"),
        "address": legal_config_value("LEGAL_POSTAL_ADDRESS"),
        "email": legal_config_value("LEGAL_CONTACT_EMAIL", "malsori91@gmail.com"),
        "phone": legal_config_value("LEGAL_CONTACT_PHONE"),
    }


def render_imprint() -> None:
    contact = legal_contact_details()
    st.markdown("### Impressum")
    st.caption("Angaben gemäß § 5 Digitale-Dienste-Gesetz (DDG)")
    if not contact["address"]:
        st.warning(
            "Vor der Veröffentlichung fehlt noch eine vollständige, ladungsfähige "
            "Postanschrift. Trage sie als LEGAL_POSTAL_ADDRESS in den Streamlit-Secrets ein."
        )

    st.markdown(f"**Diensteanbieter und Verantwortlicher**  \n{escape(contact['name'])}")
    if contact["address"]:
        address_html = "<br>".join(
            escape(line.strip())
            for line in contact["address"].splitlines()
            if line.strip()
        )
        st.markdown(address_html, unsafe_allow_html=True)
    else:
        st.markdown("*Postanschrift wird vor der öffentlichen Freigabe ergänzt.*")

    st.markdown(f"**E-Mail:** {escape(contact['email'])}")
    if contact["phone"]:
        st.markdown(f"**Telefon:** {escape(contact['phone'])}")

    st.info(
        "NursePlan Pro ist derzeit eine kostenlose Test- und Portfolioanwendung. "
        "Die Anwendung ist kein medizinisches Produkt und ersetzt keine rechtliche, "
        "tarifliche oder betriebliche Prüfung eines Dienstplans."
    )


def render_privacy_notice() -> None:
    contact = legal_contact_details()
    st.markdown("### Datenschutzerklärung")
    st.caption("Informationen gemäß Art. 13 und 14 Datenschutz-Grundverordnung (DSGVO)")

    st.markdown("#### 1. Verantwortlicher")
    st.write(contact["name"])
    if contact["address"]:
        st.write(contact["address"])
    else:
        st.warning("Die Postanschrift des Verantwortlichen muss vor der Freigabe ergänzt werden.")
    st.write(f'E-Mail: {contact["email"]}')

    st.markdown("#### 2. Welche Daten verarbeitet werden")
    st.markdown(
        """
- **Kontodaten:** Name, E-Mail-Adresse, interne Nutzer-ID und Anmeldeinformationen.
- **Planungsdaten:** Namen und Beschäftigungsdaten von Mitarbeitenden, Abwesenheiten, Planungsregeln, Vorlagen und Dienstpläne.
- **Kontaktdaten:** Inhalt einer freiwillig gesendeten Nachricht und die E-Mail-Adresse des angemeldeten Kontos.
- **Technische Daten:** IP-Adresse, Zeitpunkt, Browser- und Serverinformationen sowie technisch notwendige Sitzungsdaten.
"""
    )

    st.markdown("#### 3. Zwecke und Rechtsgrundlagen")
    st.markdown(
        """
Die Daten werden zur Registrierung, Anmeldung, Speicherung persönlicher Planungsdaten,
Erstellung von Dienstplanvorschlägen, Beantwortung von Nachrichten sowie zur sicheren
Bereitstellung der Anwendung verarbeitet. Rechtsgrundlagen sind – abhängig von der
konkreten Nutzung – Art. 6 Abs. 1 Buchst. b DSGVO (Nutzung der Anwendung), Art. 6
Abs. 1 Buchst. f DSGVO (sicherer und stabiler Betrieb) und bei freiwilligen Angaben
Art. 6 Abs. 1 Buchst. a DSGVO.
"""
    )

    st.markdown("#### 4. Eingesetzte Dienste und Empfänger")
    st.markdown(
        """
- **Streamlit Community Cloud / Snowflake:** Hosting und technische Bereitstellung der Webanwendung.
- **Supabase:** Benutzeranmeldung, Datenbank und nutzerbezogene Zugriffskontrolle.
- **Google:** nur wenn freiwillig „Mit Google anmelden“ gewählt wird. Das Google-Passwort wird NursePlan Pro nicht mitgeteilt.

Bei diesen Anbietern kann eine Verarbeitung außerhalb der Europäischen Union nicht
vollständig ausgeschlossen werden. Vor einer produktiven Nutzung mit echten
Beschäftigtendaten müssen die konkreten Standorte, Verträge zur Auftragsverarbeitung
und Garantien für Drittlandübermittlungen geprüft und dokumentiert werden.
"""
    )

    st.markdown("#### 5. Technisch notwendige Speicherung im Browser")
    st.markdown(
        """
Wenn „Auf diesem Gerät angemeldet bleiben“ aktiviert ist, speichert der Browser ein
technisch notwendiges Anmelde-Cookie für höchstens 30 Tage. Es dient ausschließlich
der Wiederherstellung der gewünschten Anmeldung. Es werden keine Werbe- oder
Tracking-Cookies durch NursePlan Pro gesetzt.
"""
    )

    st.markdown("#### 6. Speicherdauer")
    st.markdown(
        """
Konten und gespeicherte Planungsdaten bleiben grundsätzlich bis zur Löschung des
Kontos oder bis zu einem berechtigten Löschverlangen gespeichert. Kontaktanfragen
werden gelöscht, sobald sie abschließend bearbeitet wurden und keine gesetzlichen
Aufbewahrungspflichten entgegenstehen. Technische Protokolle richten sich zusätzlich
nach den Aufbewahrungsfristen der eingesetzten Hosting- und Datenbankanbieter.
"""
    )

    st.markdown("#### 7. Rechte betroffener Personen")
    st.markdown(
        f"""
Betroffene Personen können – soweit die gesetzlichen Voraussetzungen vorliegen –
Auskunft, Berichtigung, Löschung, Einschränkung der Verarbeitung und
Datenübertragbarkeit verlangen sowie einer Verarbeitung widersprechen oder eine
Einwilligung für die Zukunft widerrufen. Anfragen können an **{escape(contact['email'])}**
gerichtet werden. Außerdem besteht ein Beschwerderecht bei einer zuständigen
Datenschutzaufsichtsbehörde.
"""
    )

    st.markdown("#### 8. Automatische Dienstplanung")
    st.markdown(
        """
Die automatische Planung erzeugt nur einen Vorschlag. Sie trifft keine rechtlich
verbindliche Personalentscheidung. Jeder Plan muss vor Verwendung von einer
verantwortlichen Person fachlich, arbeitsrechtlich und betrieblich geprüft werden.
"""
    )


def render_usage_rules() -> None:
    st.markdown("### Nutzungshinweise und gesetzliche Grenzen")
    st.markdown(
        """
1. NursePlan Pro ist derzeit eine **kostenlose Test- und Portfolioanwendung**.
2. Bis zur abgeschlossenen Datenschutzprüfung dürfen nur erfundene oder vollständig anonymisierte Testdaten eingegeben werden.
3. **Patientendaten, Gesundheitsdaten, Passwörter und Zugangsschlüssel dürfen nicht eingegeben werden.**
4. Wer Daten von Mitarbeitenden eingibt, muss selbst eine zulässige Rechtsgrundlage besitzen und die betroffenen Personen informieren.
5. Automatisch erstellte Pläne müssen insbesondere auf Arbeitszeitgesetz, Tarifvertrag, Betriebsvereinbarung, Qualifikationsvorgaben, Ruhezeiten und individuelle Schutzregelungen geprüft werden.
6. Nutzer müssen ihr Konto schützen und dürfen die Anwendung nicht missbräuchlich verwenden.
7. Dienstplanvorschläge dürfen erst nach menschlicher Prüfung freigegeben werden.
"""
    )
    st.warning(
        "Vor einer betrieblichen Nutzung mit echten Beschäftigtendaten sind mindestens "
        "Datenschutzerklärung, Auftragsverarbeitungsverträge, Löschkonzept, "
        "Berechtigungskonzept und die betrieblichen Planungsregeln fachlich und rechtlich zu prüfen."
    )


def render_legal_page(selected: str = "all") -> None:
    st.markdown(
        """
        <div class="np-card" style="margin:1rem 0 1.25rem 0;">
            <div style="font-size:2rem;font-weight:800;">Rechtliches</div>
            <div style="color:#667085;margin-top:.25rem;">
                Impressum, Datenschutz und klare Grenzen der Testversion.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    if selected == "impressum":
        render_imprint()
    elif selected == "datenschutz":
        render_privacy_notice()
    elif selected == "nutzung":
        render_usage_rules()
    else:
        imprint_tab, privacy_tab, rules_tab = st.tabs(
            ["Impressum", "Datenschutz", "Nutzungshinweise"]
        )
        with imprint_tab:
            render_imprint()
        with privacy_tab:
            render_privacy_notice()
        with rules_tab:
            render_usage_rules()


def show_auth_page() -> None:
    # Supabase kann bei erlaubten Site-URLs nur zur Startseite zurückleiten.
    # Daher prüft auch die normale Anmeldeseite auf ein OAuth-URL-Fragment.
    show_google_oauth_callback(show_status=False)

    st.markdown(
        """
        <style>
        header[data-testid="stHeader"] {
            background: transparent !important;
            border-bottom: 0 !important;
        }
        [data-testid="stMainBlockContainer"] {
            max-width: 1220px;
            padding-top: 3.4rem;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand) {
            gap: 0 !important;
            align-items: stretch !important;
            border-radius: 24px;
            box-shadow: 0 24px 70px rgba(15, 45, 70, .14);
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) {
            background: #ffffff;
            border-radius: 0 24px 24px 0;
            padding: 2.4rem 3.1rem 2rem 3.1rem;
        }
        .np-auth-brand {
            position: relative;
            min-height: 720px;
            height: 100%;
            overflow: hidden;
            border-radius: 24px 0 0 24px;
            padding: 4.6rem 3.6rem;
            color: #ffffff;
            background:
                radial-gradient(circle at 78% 20%, rgba(59,183,160,.15), transparent 28%),
                linear-gradient(145deg, #0b2944 0%, #0f3658 100%);
        }
        .np-auth-brand::after {
            content: "";
            position: absolute;
            width: 590px;
            height: 360px;
            right: -150px;
            bottom: -95px;
            opacity: .11;
            transform: rotate(-12deg);
            background-image:
                linear-gradient(rgba(255,255,255,.65) 1px, transparent 1px),
                linear-gradient(90deg, rgba(255,255,255,.65) 1px, transparent 1px);
            background-size: 82px 64px;
        }
        .np-auth-brand-content { position: relative; z-index: 1; }
        .np-auth-logo {
            width: 70px;
            height: 70px;
            display: grid;
            place-items: center;
            margin-bottom: 1.45rem;
            border: 2px solid rgba(255,255,255,.75);
            border-radius: 18px;
            font-size: 2rem;
            box-shadow: inset 0 0 0 6px rgba(255,255,255,.05);
        }
        .np-auth-brand h1 {
            margin: 0;
            color: #ffffff;
            font-size: 3.2rem;
            line-height: 1.04;
            letter-spacing: -.055em;
            font-weight: 850;
        }
        .np-auth-brand h1 span { color: #3bb7a0; }
        .np-auth-tagline {
            margin-top: 1rem;
            color: rgba(255,255,255,.88);
            font-size: 1.18rem;
            line-height: 1.55;
        }
        .np-auth-features {
            display: grid;
            gap: 1.35rem;
            margin-top: 3.1rem;
        }
        .np-auth-feature {
            display: flex;
            align-items: center;
            gap: .9rem;
            color: rgba(255,255,255,.92);
            font-size: 1rem;
            font-weight: 600;
        }
        .np-auth-feature-icon {
            width: 38px;
            height: 38px;
            display: grid;
            place-items: center;
            flex: 0 0 38px;
            border: 1px solid rgba(255,255,255,.35);
            border-radius: 11px;
            color: #55d0bb;
            background: rgba(255,255,255,.06);
        }
        .np-auth-card-head h2 {
            margin: 0;
            color: #0f2d46;
            font-size: 2.15rem;
            line-height: 1.12;
            letter-spacing: -.04em;
            font-weight: 850;
        }
        .np-auth-card-head p {
            margin: .55rem 0 1.55rem 0;
            color: #667085;
            font-size: 1rem;
        }
        .np-auth-divider {
            display: flex;
            align-items: center;
            gap: .85rem;
            margin: 1.15rem 0 1rem 0;
            color: #667085;
            font-size: .9rem;
        }
        .np-auth-divider::before,
        .np-auth-divider::after {
            content: "";
            height: 1px;
            flex: 1;
            background: #dfe5ec;
        }
        .np-auth-security {
            margin-top: 1.1rem;
            padding: .78rem .9rem;
            border-radius: 12px;
            color: #526071;
            background: #f5f8fb;
            font-size: .8rem;
            line-height: 1.45;
        }
        .np-auth-footer {
            padding: 1.25rem 0 .2rem 0;
            text-align: center;
            color: #667085;
            font-size: .88rem;
        }
        .np-auth-footer a {
            color: #526071;
            text-decoration: none;
            margin: 0 .35rem;
        }
        .np-auth-footer a:hover { color: #0f8e80; text-decoration: underline; }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) div[data-testid="stForm"] {
            padding: 0;
            border: 0;
            border-radius: 0;
            box-shadow: none;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) div[data-baseweb="input"] {
            min-height: 3.15rem;
            border-color: #ccd6e1 !important;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) button[kind="primary"] {
            min-height: 3.15rem;
            background: #ff5258 !important;
            border-color: #ff5258 !important;
            box-shadow: 0 8px 18px rgba(255,82,88,.2);
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) button[kind="primary"]:hover {
            background: #ef444a !important;
            border-color: #ef444a !important;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) [data-testid="stLinkButton"] a {
            min-height: 3.15rem;
            border: 1.5px solid #163d5f;
            border-radius: 11px;
            color: #0f2d46;
            background: #ffffff;
            font-weight: 750;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) [data-testid="stLinkButton"] a:hover {
            border-color: #0f8e80;
            color: #0f8e80;
            background: #f6fffd;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) [data-testid="stLinkButton"] a::before {
            content: "G";
            margin-right: .48rem;
            color: #4285f4;
            font-size: 1.15rem;
            font-weight: 900;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) .stButton button[kind="secondary"] {
            border: 0;
            color: #0f8e80;
            background: transparent;
            box-shadow: none;
        }
        div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
        > div[data-testid="stColumn"]:has(.np-auth-card-head) .stButton button[kind="secondary"]:hover {
            color: #0b6f65;
            background: #f0fbf9;
        }
        @media (max-width: 768px) {
            [data-testid="stMainBlockContainer"] {
                padding: .8rem .7rem 2rem .7rem;
            }
            div[data-testid="stHorizontalBlock"]:has(.np-auth-brand) {
                display: flex !important;
                flex-direction: column !important;
                box-shadow: 0 16px 42px rgba(15,45,70,.12);
            }
            div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
            > div[data-testid="stColumn"] {
                width: 100% !important;
                min-width: 100% !important;
                flex: 1 1 100% !important;
            }
            .np-auth-brand {
                min-height: 0;
                border-radius: 22px 22px 0 0;
                padding: 2rem 1.5rem 2.15rem 1.5rem;
                text-align: center;
            }
            .np-auth-brand::after { opacity: .07; }
            .np-auth-logo {
                width: 58px;
                height: 58px;
                margin: 0 auto 1rem auto;
                font-size: 1.6rem;
            }
            .np-auth-brand h1 { font-size: 2.35rem; }
            .np-auth-tagline { margin: .65rem auto 0 auto; font-size: 1rem; }
            .np-auth-features { display: none; }
            div[data-testid="stHorizontalBlock"]:has(.np-auth-brand)
            > div[data-testid="stColumn"]:has(.np-auth-card-head) {
                border-radius: 0 0 22px 22px;
                padding: 1.75rem 1.25rem 1.45rem 1.25rem;
            }
            .np-auth-card-head h2 { font-size: 1.8rem; }
            .np-auth-card-head p { margin-bottom: 1.2rem; }
            .np-auth-footer { font-size: .78rem; }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    password_reset_success = st.session_state.pop("password_reset_success", False)
    google_oauth_error = st.session_state.pop("google_oauth_error", False)
    google_session_error = st.session_state.pop("google_oauth_session_error", "")
    reset_mode = bool(st.session_state.get("auth_show_password_reset"))
    auth_mode = str(st.session_state.get("auth_mode", "login"))
    if auth_mode not in {"login", "register"}:
        auth_mode = "login"

    brand_column, form_column = st.columns([1.05, 1], gap="small")

    with brand_column:
        st.markdown(
            """
            <div class="np-auth-brand">
                <div class="np-auth-brand-content">
                    <div class="np-auth-logo">▦<span style="color:#55d0bb;">＋</span></div>
                    <h1>NursePlan <span>Pro</span></h1>
                    <div class="np-auth-tagline">Dienstplanung, die zu deinem Team passt.</div>
                    <div class="np-auth-features">
                        <div class="np-auth-feature">
                            <div class="np-auth-feature-icon">✓</div>
                            <div>Sicher und geschützt</div>
                        </div>
                        <div class="np-auth-feature">
                            <div class="np-auth-feature-icon">▦</div>
                            <div>Einfach planen</div>
                        </div>
                        <div class="np-auth-feature">
                            <div class="np-auth-feature-icon">⌁</div>
                            <div>Überall verfügbar</div>
                        </div>
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with form_column:
        if reset_mode:
            heading = "Passwort zurücksetzen"
            subheading = "Wir senden dir einen sicheren Link an deine E-Mail-Adresse."
        elif auth_mode == "register":
            heading = "Konto erstellen"
            subheading = "Starte kostenlos und plane dein Team übersichtlich."
        else:
            heading = "Willkommen zurück"
            subheading = "Melde dich an und plane dein Team."

        st.markdown(
            f"""
            <div class="np-auth-card-head">
                <h2>{escape(heading)}</h2>
                <p>{escape(subheading)}</p>
            </div>
            """,
            unsafe_allow_html=True,
        )

        if password_reset_success:
            st.success("Dein Passwort wurde geändert. Du kannst dich jetzt anmelden.")
        if google_oauth_error:
            st.error(
                "Die Google-Anmeldung wurde abgebrochen oder konnte nicht abgeschlossen werden."
            )
        if google_session_error:
            st.error(google_session_error)

        if reset_mode:
            if st.session_state.get("auth_reset_email_sent"):
                st.success(
                    "Wenn für diese Adresse ein Konto besteht, wurde eine E-Mail gesendet. "
                    "Bitte prüfe auch den Spam-Ordner."
                )
            else:
                with st.form("password_reset_request_form"):
                    reset_email = st.text_input(
                        "E-Mail-Adresse",
                        key="password_reset_email",
                        placeholder="name@beispiel.de",
                        autocomplete="email",
                    )
                    submit_reset = st.form_submit_button(
                        "Link senden", type="primary", width="stretch"
                    )
                if submit_reset:
                    if not reset_email.strip() or "@" not in reset_email:
                        st.warning("Bitte gib eine gültige E-Mail-Adresse ein.")
                    else:
                        try:
                            request_password_reset(
                                reset_email,
                                password_reset_redirect_url(),
                            )
                            st.session_state.auth_reset_email_sent = True
                            st.rerun()
                        except Exception as exc:
                            st.error(auth_error_text(exc, "password_reset"))

            if st.button(
                "← Zurück zur Anmeldung",
                type="secondary",
                width="stretch",
                key="auth_back_from_reset",
            ):
                st.session_state.auth_show_password_reset = False
                st.session_state.pop("auth_reset_email_sent", None)
                st.rerun()

            st.markdown(
                '<div class="np-auth-security">🔒 Der Link kann nur einmal verwendet werden und läuft nach kurzer Zeit ab.</div>',
                unsafe_allow_html=True,
            )

        else:
            try:
                google_url = google_oauth_url()
            except Exception:
                google_url = ""

            if google_url:
                st.link_button(
                    "Mit Google anmelden",
                    google_url,
                    type="secondary",
                    width="stretch",
                )
                st.markdown(
                    '<div class="np-auth-divider">oder mit E-Mail</div>',
                    unsafe_allow_html=True,
                )
            else:
                st.warning("Die Google-Anmeldung ist vorübergehend nicht verfügbar.")

            if auth_mode == "register":
                with st.form("register_form"):
                    full_name = st.text_input(
                        "Vollständiger Name",
                        key="register_name",
                        placeholder="Vorname Nachname",
                        autocomplete="name",
                    )
                    register_email = st.text_input(
                        "E-Mail-Adresse",
                        key="register_email",
                        placeholder="name@beispiel.de",
                        autocomplete="email",
                    )
                    register_password = st.text_input(
                        "Passwort (mindestens 8 Zeichen)",
                        type="password",
                        key="register_password",
                        autocomplete="new-password",
                    )
                    password_repeat = st.text_input(
                        "Passwort bestätigen",
                        type="password",
                        key="register_password_repeat",
                        autocomplete="new-password",
                    )
                    remember_registration = st.checkbox(
                        "Angemeldet bleiben",
                        value=True,
                        key="register_remember_login",
                        help="Nur auf einem persönlichen oder geschützten Gerät verwenden.",
                    )
                    submit_registration = st.form_submit_button(
                        "Kostenlos registrieren", type="primary", width="stretch"
                    )
                if submit_registration:
                    if not full_name.strip() or not register_email.strip():
                        st.warning("Bitte Name und E-Mail-Adresse eingeben.")
                    elif len(register_password) < 8:
                        st.warning("Das Passwort muss mindestens 8 Zeichen haben.")
                    elif register_password != password_repeat:
                        st.warning("Die beiden Passwörter stimmen nicht überein.")
                    else:
                        try:
                            refresh_token = sign_up_user(
                                full_name, register_email, register_password
                            )
                            if refresh_token:
                                st.session_state.pop("auth_logout_requested", None)
                                st.session_state.pop("auth_restore_failed", None)
                                st.session_state.remember_login = bool(
                                    remember_registration
                                )
                                if remember_registration:
                                    save_auth_cookie(refresh_token)
                                else:
                                    delete_auth_cookie()
                                st.rerun()
                            else:
                                st.success(
                                    "Das Konto wurde angelegt. Öffne jetzt die "
                                    "Bestätigungs-E-Mail und melde dich danach an."
                                )
                        except Exception as exc:
                            st.error(auth_error_text(exc, "registration"))

                if st.button(
                    "Bereits registriert? Zur Anmeldung",
                    type="secondary",
                    width="stretch",
                    key="auth_switch_to_login",
                ):
                    st.session_state.auth_mode = "login"
                    st.rerun()

            else:
                with st.form("login_form"):
                    email = st.text_input(
                        "E-Mail-Adresse",
                        key="login_email",
                        placeholder="name@beispiel.de",
                        autocomplete="email",
                    )
                    password = st.text_input(
                        "Passwort",
                        type="password",
                        key="login_password",
                        autocomplete="current-password",
                    )
                    remember_login = st.checkbox(
                        "Angemeldet bleiben",
                        value=True,
                        help="Nur auf einem persönlichen oder geschützten Gerät verwenden.",
                    )
                    submit_login = st.form_submit_button(
                        "Sicher anmelden", type="primary", width="stretch"
                    )
                if submit_login:
                    if not email.strip() or not password:
                        st.warning("Bitte E-Mail-Adresse und Passwort eingeben.")
                    else:
                        try:
                            refresh_token = sign_in_user(email, password)
                            st.session_state.pop("auth_logout_requested", None)
                            st.session_state.pop("auth_restore_failed", None)
                            st.session_state.remember_login = bool(remember_login)
                            if remember_login:
                                save_auth_cookie(refresh_token)
                            else:
                                delete_auth_cookie()
                            st.rerun()
                        except Exception as exc:
                            st.error(auth_error_text(exc, "login"))

                if st.button(
                    "Passwort vergessen?",
                    type="secondary",
                    width="stretch",
                    key="auth_show_reset",
                ):
                    st.session_state.auth_show_password_reset = True
                    st.session_state.pop("auth_reset_email_sent", None)
                    st.rerun()

                if st.button(
                    "Noch kein Konto? Kostenlos registrieren",
                    type="secondary",
                    width="stretch",
                    key="auth_switch_to_register",
                ):
                    st.session_state.auth_mode = "register"
                    st.rerun()

            st.markdown(
                '<div class="np-auth-security">🔒 Deine Mitarbeitenden, Abwesenheiten und Dienstpläne bleiben deinem geschützten Konto zugeordnet.</div>',
                unsafe_allow_html=True,
            )

    st.markdown(
        """
        <div class="np-auth-footer">
            <a href="?legal=impressum" target="_self">Impressum</a> ·
            <a href="?legal=datenschutz" target="_self">Datenschutz</a> ·
            <a href="?legal=nutzung" target="_self">Nutzungshinweise</a>
        </div>
        """,
        unsafe_allow_html=True,
    )


try:
    get_supabase_client()
except SupabaseConfigurationError as exc:
    st.error("Supabase ist noch nicht mit der App verbunden.")
    st.write(str(exc))
    st.code(
        'SUPABASE_URL = "https://DEIN-PROJEKT.supabase.co"\n'
        'SUPABASE_PUBLISHABLE_KEY = "DEIN-PUBLISHABLE-KEY"',
        language="toml",
    )
    st.caption(
        "Diese zwei Werte gehören lokal in .streamlit/secrets.toml. "
        "Niemals den Secret- oder service_role-Schlüssel verwenden."
    )
    st.stop()

public_legal_page = str(st.query_params.get("legal", "")).casefold()
if public_legal_page in {"impressum", "datenschutz", "nutzung"}:
    render_legal_page(public_legal_page)
    st.markdown("[← Zurück zur Anmeldung](?)")
    st.stop()

# Nach der Rückkehr von Google liegt die Supabase-Sitzung kurz im URL-Fragment.
# Ein kleines Browser-Hilfsfenster speichert die neue Sitzung kurzzeitig im
# NursePlan-Cookie und entfernt anschließend sämtliche Token aus der Adresse.
if str(st.query_params.get("oauth_callback", "")).casefold() == "google":
    show_google_oauth_callback()
    st.stop()

if str(st.query_params.get("oauth_complete", "")).casefold() == "1":
    st.session_state.pop("auth_logout_requested", None)
    st.session_state.pop("auth_restore_failed", None)
    st.query_params.clear()

if str(st.query_params.get("oauth_error", "")).casefold() == "1":
    st.session_state.google_oauth_error = True
    st.query_params.clear()

# Nach einer erfolgreichen Änderung kehrt der Browser ohne Sitzung zur Anmeldung zurück.
if str(st.query_params.get("password_reset", "")).casefold() == "success":
    delete_auth_cookie()
    clear_local_auth_state()
    st.session_state.auth_logout_requested = True
    st.session_state.password_reset_success = True
    st.query_params.clear()

# Der Supabase-Link enthält die kurzlebige Wiederherstellungs-Sitzung im URL-Fragment.
# Das Passwortformular verarbeitet sie direkt im Browser, damit kein Passwort die App erreicht.
if str(st.query_params.get("password_recovery", "")).casefold() == "1":
    show_password_recovery_page()
    st.stop()

# Eine gespeicherte Supabase-Sitzung wird nach Browser-Refresh erneuert.
# Refresh-Tokens werden dabei rotiert und sofort wieder im Cookie aktualisiert.
if (
    not is_authenticated()
    and not st.session_state.get("auth_logout_requested")
    and not st.session_state.get("auth_restore_failed")
):
    stored_access_token, stored_refresh_token = saved_auth_tokens()
    if stored_refresh_token:
        try:
            if stored_access_token:
                rotated_refresh_token = adopt_google_user_session(
                    stored_access_token,
                    stored_refresh_token,
                )
            else:
                rotated_refresh_token = restore_user_session(stored_refresh_token)
            st.session_state.remember_login = True
            save_auth_cookie(rotated_refresh_token)
        except Exception as exc:
            clear_local_auth_state()
            # Ein abgelaufener oder bereits rotierter Refresh-Token darf nicht
            # bei jedem neuen Seitenaufruf erneut an Supabase gesendet werden.
            delete_auth_cookie()
            st.session_state.auth_restore_failed = True
            if stored_access_token:
                st.session_state.google_oauth_session_error = (
                    google_oauth_session_error_text(exc)
                )

if is_authenticated() and st.session_state.get("remember_login"):
    current_refresh_token = current_session_refresh_token()
    if current_refresh_token:
        save_auth_cookie(current_refresh_token)

if not is_authenticated():
    show_auth_page()
    st.stop()

create_tables()

if "settings" not in st.session_state:
    raw_settings = get_einstellungen(DEFAULT_SETTINGS)
    # Werte aus älteren Versionen übernehmen.
    if "n_start" not in raw_settings and "na_start" in raw_settings:
        raw_settings["n_start"] = raw_settings.get("na_start", "21:48")
        raw_settings["n_ende"] = raw_settings.get("na_ende", "06:00")
        raw_settings["n_stunden"] = raw_settings.get("na_stunden", 7.7)
        raw_settings["min_n"] = raw_settings.get("min_na", 1)
    st.session_state.settings = {**DEFAULT_SETTINGS, **raw_settings}

if "dienstplan_vorlage" not in st.session_state:
    st.session_state.dienstplan_vorlage = []
if "dienstplan_editor_version" not in st.session_state:
    st.session_state.dienstplan_editor_version = 0
if "aktiver_dienstplan_id" not in st.session_state:
    st.session_state.aktiver_dienstplan_id = None
if "dienstplan_name" not in st.session_state:
    st.session_state.dienstplan_name = "Neuer Dienstplan"
if "auto_plan_result" not in st.session_state:
    st.session_state.auto_plan_result = None


def page_header(title: str, subtitle: str) -> None:
    st.markdown(f'<p class="np-title">{escape(title)}</p>', unsafe_allow_html=True)
    st.markdown(f'<p class="np-subtitle">{escape(subtitle)}</p>', unsafe_allow_html=True)


@lru_cache(maxsize=128)
def parse_clock(value: str) -> time:
    try:
        return datetime.strptime(str(value), "%H:%M").time()
    except (TypeError, ValueError):
        return time(0, 0)


def format_clock(value: time) -> str:
    return value.strftime("%H:%M")


def percentage_from_hours(monthly_hours: float) -> float:
    basis = float(st.session_state.settings.get("vollzeit_monatssoll", 167.0))
    if basis <= 0:
        return 0.0
    return round(float(monthly_hours) / basis * 100.0, 1)


def percentage_text(monthly_hours: float) -> str:
    value = percentage_from_hours(monthly_hours)
    if value.is_integer():
        return f"{int(value)} %"
    return f"{value:.1f} %".replace(".", ",")


def weekly_from_month(monthly_hours: float) -> float:
    return round(float(monthly_hours) * 12.0 / 52.0, 2)


def stored_percentage(monthly_hours: float) -> str:
    return f"{percentage_from_hours(monthly_hours):.1f}%"


def monthly_target(person: dict) -> float:
    try:
        return round(float(person.get("monatssollstunden", 0.0)), 2)
    except (TypeError, ValueError):
        return 0.0


def day_column(day_value: date) -> str:
    return f"{day_value.day:02d} {WOCHENTAGE[day_value.weekday()]}"


def split_names(value) -> list[str]:
    if value is None or pd.isna(value):
        return []
    return [x.strip() for x in re.split(r"[,;\n]+", str(value)) if x.strip()]


def join_names(values) -> str:
    result = []
    seen = set()
    for value in values:
        name = str(value).strip()
        key = name.casefold()
        if name and key not in seen:
            result.append(name)
            seen.add(key)
    return ", ".join(result)


@lru_cache(maxsize=1024)
def parse_date(value):
    try:
        return datetime.strptime(str(value), "%d.%m.%Y").date()
    except (TypeError, ValueError):
        return None


def weekend_anchor(day_value: date | None) -> date | None:
    """Gibt den Samstag eines Wochenendes zurück; Samstag und Sonntag gehören zusammen."""
    if day_value is None:
        return None
    if day_value.weekday() == 5:
        return day_value
    if day_value.weekday() == 6:
        return day_value - timedelta(days=1)
    return None


def worked_weekends(values) -> set[date]:
    return {anchor for value in values if (anchor := weekend_anchor(value)) is not None}


def has_consecutive_weekends(values: set[date]) -> bool:
    ordered = sorted(values)
    return any(current == previous + timedelta(days=7) for previous, current in zip(ordered, ordered[1:]))


def regel_aktiv(settings: dict, key: str, default: bool = True) -> bool:
    return bool(settings.get(key, default))


def eigene_planungsregeln(settings: dict, regeltyp: str | None = None) -> list[dict]:
    rules = settings.get("eigene_planungsregeln", [])
    if not isinstance(rules, list):
        return []
    result = [dict(rule) for rule in rules if isinstance(rule, dict)]
    if regeltyp is not None:
        result = [rule for rule in result if str(rule.get("typ")) == regeltyp]
    return result


def regel_gilt_fuer_person(rule: dict, person: dict) -> bool:
    try:
        person_id = int(rule.get("mitarbeitende_id", 0))
        return person_id == 0 or person_id == int(person.get("id", -1))
    except (TypeError, ValueError):
        return False


def regel_zahl(rule: dict, key: str, default: int) -> int:
    try:
        return int(rule.get(key, default))
    except (TypeError, ValueError):
        return int(default)


def regel_wochentage(rule: dict) -> set[int]:
    """Liest gespeicherte Wochentage sicher als Werte von Montag=0 bis Sonntag=6."""
    values = rule.get("wochentage", [])
    if not isinstance(values, list):
        return set()
    result = set()
    for value in values:
        try:
            weekday = int(value)
        except (TypeError, ValueError):
            continue
        if 0 <= weekday <= 6:
            result.add(weekday)
    return result


def mitarbeitende_name(employee_id, employees) -> str:
    try:
        target = int(employee_id)
    except (TypeError, ValueError):
        return "Unbekannt"
    if target == 0:
        return "Alle Mitarbeitenden"
    for person in employees:
        if int(person.get("id", -1)) == target:
            return str(person.get("name", "Unbekannt"))
    return "Gelöschte Person"


def regel_beschreibung(rule: dict, employees) -> str:
    shift_names = {"F": "Frühdienst", "S": "Spätdienst", "N": "Nachtdienst", "ALL": "allen Diensten"}
    rule_type = str(rule.get("typ", ""))
    if rule_type == "nicht_gemeinsam":
        first = mitarbeitende_name(rule.get("mitarbeitende_a_id"), employees)
        second = mitarbeitende_name(rule.get("mitarbeitende_b_id"), employees)
        code = str(rule.get("dienst", "ALL"))
        if code == "ALL":
            return f"{first} und {second} dürfen nicht gemeinsam in derselben Schicht arbeiten."
        shift = shift_names.get(code, "gewählten Dienst")
        return f"{first} und {second} dürfen nicht gemeinsam im {shift} eingeplant werden."
    if rule_type == "frei_nach_dienst":
        person = mitarbeitende_name(rule.get("mitarbeitende_id", 0), employees)
        shift = shift_names.get(str(rule.get("dienst", "N")), "Dienst")
        days = regel_zahl(rule, "freie_tage", 1)
        return f"{person}: Nach einem {shift} folgen mindestens {days} freie Kalendertage."
    if rule_type == "max_dienstfolge":
        person = mitarbeitende_name(rule.get("mitarbeitende_id", 0), employees)
        shift = shift_names.get(str(rule.get("dienst", "N")), "Dienst")
        maximum = regel_zahl(rule, "maximum", 1)
        return f"{person}: Höchstens {maximum} {shift}-Dienste hintereinander."
    if rule_type == "nicht_an_wochentagen":
        person = mitarbeitende_name(rule.get("mitarbeitende_id", 0), employees)
        weekdays = [WOCHENTAGE[value] for value in sorted(regel_wochentage(rule))]
        day_text = ", ".join(weekdays) if weekdays else "keinen ausgewählten Tagen"
        return f"{person}: Keine Arbeit an folgenden Wochentagen: {day_text}."
    return "Unbekannte Regelart"


def normalise_plan(plan):
    result = []
    for index, entry in enumerate(plan or [], start=1):
        row = dict(entry)
        row.setdefault("Tag", index)
        row.setdefault("Datum", "")
        row.setdefault("Frühdienst", "")
        row.setdefault("Spätdienst", "")
        row.setdefault("Nachtdienst", "")
        row.setdefault("Urlaub", "")
        row.setdefault("Wunschfrei", "")
        row.setdefault("Fortbildung", "")
        row.setdefault("Krank", "")
        row.setdefault("Frei", "")
        row.setdefault("Status", "")
        result.append(row)
    return result

def create_empty_month(year: int, month: int):
    return [
        {
            "Tag": day,
            "Datum": date(year, month, day).strftime("%d.%m.%Y"),
            "Frühdienst": "",
            "Spätdienst": "",
            "Nachtdienst": "",
            "Urlaub": "",
            "Wunschfrei": "",
            "Fortbildung": "",
            "Krank": "",
            "Frei": "",
            "Status": "",
        }
        for day in range(1, calendar.monthrange(year, month)[1] + 1)
    ]


FACHKRAFT_QUALIFIKATIONEN = {
    "Pflegefachkraft",
    "Praxisanleitung",
    "Stroke Nurse",
    "Stationsleitung",
}

VORGABE_LABELS = {
    "U": "Urlaub",
    "W": "Wunschfrei",
    "FB": "Fortbildung",
    "K": "Krank",
}

# Alte Versionen erlaubten F/S/N/Frei als datumsbezogene Vorgabe.
# Diese Einträge werden nicht mehr automatisch angewendet, können aber in der
# Übersicht noch erkannt und gelöscht werden.
LEGACY_VORGABE_LABELS = {
    "F": "Alte Vorgabe: Frühdienst (wird ignoriert)",
    "S": "Alte Vorgabe: Spätdienst (wird ignoriert)",
    "N": "Alte Vorgabe: Nachtdienst (wird ignoriert)",
    "Frei": "Alte Vorgabe: fest frei (wird ignoriert)",
}

CODE_TO_COLUMN = {
    "F": "Frühdienst",
    "S": "Spätdienst",
    "N": "Nachtdienst",
    "U": "Urlaub",
    "W": "Wunschfrei",
    "FB": "Fortbildung",
    "K": "Krank",
    "Frei": "Frei",
}


def is_fachkraft(person: dict) -> bool:
    return str(person.get("qualifikation", "")) in FACHKRAFT_QUALIFIKATIONEN


def erlaubte_dienst_codes(person: dict) -> set[str]:
    feste_dienstart = str(person.get("feste_dienstart", "Alle Dienste")).strip()
    return FESTE_DIENSTART_CODES.get(feste_dienstart, {"F", "S", "N"})


def feste_dienstart_text(person: dict) -> str:
    value = str(person.get("feste_dienstart", "Alle Dienste")).strip()
    return value if value in FESTE_DIENSTARTEN else "Alle Dienste"


def daterange(start: date, end: date):
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def remove_person_from_day(entry: dict, name: str) -> None:
    target = name.casefold()
    for column in CODE_TO_COLUMN.values():
        entry[column] = join_names(
            value
            for value in split_names(entry.get(column, ""))
            if value.casefold() != target
        )


def apply_vorgaben(year: int, month: int, employees, vorgaben):
    plan = create_empty_month(year, month)
    by_day = {parse_date(row["Datum"]): row for row in plan}
    employee_by_id = {int(person["id"]): person for person in employees}
    month_start = date(year, month, 1)
    month_end = date(year, month, calendar.monthrange(year, month)[1])

    # Neuere Vorgaben haben Vorrang, falls versehentlich mehrere Vorgaben kollidieren.
    for vorgabe in sorted(vorgaben or [], key=lambda item: int(item.get("id", 0))):
        person = employee_by_id.get(int(vorgabe.get("mitarbeitende_id", -1)))
        code = str(vorgabe.get("code", "")).strip()
        if not person or code not in VORGABE_LABELS:
            continue
        try:
            start = datetime.strptime(str(vorgabe["startdatum"]), "%Y-%m-%d").date()
            end = datetime.strptime(str(vorgabe["enddatum"]), "%Y-%m-%d").date()
        except (KeyError, TypeError, ValueError):
            continue
        start = max(start, month_start)
        end = min(end, month_end)
        if start > end:
            continue
        for day_value in daterange(start, end):
            # Urlaub wird nur an regulären Werktagen eingetragen. Ein Zeitraum
            # darf das Wochenende enthalten, Samstag und Sonntag bleiben im
            # Dienstplan jedoch frei und zählen nicht als Urlaubstage.
            if code == "U" and day_value.weekday() >= 5:
                continue
            entry = by_day.get(day_value)
            if not entry:
                continue
            remove_person_from_day(entry, person["name"])
            column = CODE_TO_COLUMN[code]
            entry[column] = join_names(split_names(entry.get(column, "")) + [person["name"]])
            entry["Status"] = "Abwesenheiten übernommen"
    return plan


def count_absence_days(plan) -> int:
    """Zählt die im Monatsplan übernommenen persönlichen Abwesenheitstage."""
    return sum(
        len(split_names(entry.get(column, "")))
        for entry in normalise_plan(plan)
        for column in ("Urlaub", "Wunschfrei", "Fortbildung", "Krank")
    )

def month_from_plan(plan):
    for entry in normalise_plan(plan):
        parsed = parse_date(entry.get("Datum"))
        if parsed:
            return parsed.year, parsed.month
    today = date.today()
    return today.year, today.month


def code_for_person_day(entry, name: str) -> str:
    target = name.casefold()
    mapping = {
        "Frühdienst": "F",
        "Spätdienst": "S",
        "Nachtdienst": "N",
        "Urlaub": "U",
        "Wunschfrei": "W",
        "Fortbildung": "FB",
        "Krank": "K",
        "Frei": "Frei",
    }
    for column, code in mapping.items():
        if any(x.casefold() == target for x in split_names(entry.get(column, ""))):
            return code
    return ""

def workdays_in_month(plan) -> int:
    year, month = month_from_plan(plan)
    return max(
        1,
        sum(
            1
            for day in range(1, calendar.monthrange(year, month)[1] + 1)
            if date(year, month, day).weekday() < 5
        ),
    )


def calculate_hours(plan, employees, settings) -> pd.DataFrame:
    plan = normalise_plan(plan)
    workdays = workdays_in_month(plan)
    rows = []

    for person in employees:
        name = person["name"]
        target = monthly_target(person)
        absence_day_hours = target / workdays if workdays else 0.0
        counts = {"F": 0, "S": 0, "N": 0, "U": 0, "W": 0, "FB": 0, "K": 0, "Frei": 0}
        credited_absence_days = 0

        for entry in plan:
            code = code_for_person_day(entry, name)
            if code in counts:
                counts[code] += 1
            day_value = parse_date(entry.get("Datum"))
            if code in {"U", "FB", "K"} and day_value and day_value.weekday() < 5:
                credited_absence_days += 1

        dienststunden = (
            counts["F"] * float(settings["f_stunden"])
            + counts["S"] * float(settings["s_stunden"])
            + counts["N"] * float(settings["n_stunden"])
        )
        abwesenheitsstunden = credited_absence_days * absence_day_hours
        actual = round(dienststunden + abwesenheitsstunden, 1)

        rows.append(
            {
                "Mitarbeitende": name,
                "Soll": round(target, 1),
                "Ist": actual,
                "+/-": round(actual - target, 1),
                "F": counts["F"],
                "S": counts["S"],
                "N": counts["N"],
                "U": counts["U"],
                "W": counts["W"],
                "FB": counts["FB"],
                "K": counts["K"],
            }
        )

    return pd.DataFrame(rows)

def plan_to_matrix(plan, employees, settings) -> pd.DataFrame:
    """Monatstabelle mit Stunden direkt neben dem Namen.

    Die Reihenfolge ist bewusst: Mitarbeitende, Soll, Ist, Rest und danach
    die einzelnen Kalendertage. So bleibt die Stundenkontrolle immer sichtbar.
    """
    plan = normalise_plan(plan)
    stats = calculate_hours(plan, employees, settings)
    stats_by_name = {
        str(row["Mitarbeitende"]).casefold(): row
        for _, row in stats.iterrows()
    }
    rows = []

    for person in employees:
        name = person["name"]
        stat = stats_by_name.get(name.casefold(), {})
        target = float(stat.get("Soll", monthly_target(person)))
        actual = float(stat.get("Ist", 0.0))
        row = {
            "Mitarbeitende": name,
            "Soll": target,
            "Ist": actual,
            "Rest": round(target - actual, 1),
        }
        for entry in plan:
            parsed = parse_date(entry.get("Datum"))
            if parsed:
                row[day_column(parsed)] = code_for_person_day(entry, name)
        rows.append(row)

    return pd.DataFrame(rows)


def matrix_to_plan(matrix: pd.DataFrame, old_plan):
    old_plan = normalise_plan(old_plan)
    result = []

    code_map = {
        "F": "Frühdienst",
        "S": "Spätdienst",
        "N": "Nachtdienst",
        "NA": "Nachtdienst",
        "U": "Urlaub",
        "W": "Wunschfrei",
        "FB": "Fortbildung",
        "K": "Krank",
        "FREI": "Frei",
    }

    for entry in old_plan:
        parsed = parse_date(entry.get("Datum"))
        if not parsed:
            continue
        column = day_column(parsed)
        assignments = {value: [] for value in CODE_TO_COLUMN.values()}
        for _, row in matrix.iterrows():
            name = str(row.get("Mitarbeitende", "")).strip()
            code = str(row.get(column, "") or "").strip().upper()
            target_column = code_map.get(code)
            if name and target_column:
                assignments[target_column].append(name)

        result.append(
            {
                "Tag": parsed.day,
                "Datum": parsed.strftime("%d.%m.%Y"),
                **{key: join_names(value) for key, value in assignments.items()},
                "Status": "Manuell bearbeitet",
            }
        )
    return result


def manual_month_input_errors(matrix: pd.DataFrame, current_plan, employees) -> list[str]:
    """Prüft die schnelle Monatseingabe, bevor sie übernommen wird."""
    people_by_name = {
        str(person.get("name", "")).strip().casefold(): person
        for person in employees
    }
    plan_by_column = {}
    for entry in normalise_plan(current_plan):
        parsed = parse_date(entry.get("Datum"))
        if parsed:
            plan_by_column[day_column(parsed)] = entry

    day_columns = [
        column for column in matrix.columns
        if re.match(r"^\d{2} ", str(column))
    ]
    protected_codes = {"U", "W", "FB", "K"}
    valid_codes = {"", "F", "S", "N", "FREI"} | protected_codes
    errors = []
    seen_errors = set()

    for _, row in matrix.iterrows():
        name = str(row.get("Mitarbeitende", "")).strip()
        person = people_by_name.get(name.casefold())
        if not person:
            continue

        for column in day_columns:
            value = row.get(column, "")
            code = "" if value is None or pd.isna(value) else str(value).strip().upper()
            if code not in valid_codes:
                message = f'{name}, {column}: „{value}“ ist kein gültiger Dienst.'
            else:
                old_entry = plan_by_column.get(column, {})
                old_code = code_for_person_day(old_entry, name).upper()
                if old_code in protected_codes and code != old_code:
                    message = (
                        f"{name}, {column}: Eine gespeicherte Abwesenheit kann hier "
                        "nicht geändert werden."
                    )
                elif old_code not in protected_codes and code in protected_codes:
                    message = (
                        f"{name}, {column}: Abwesenheiten bitte im Bereich „Abwesenheiten“ eintragen."
                    )
                elif code in {"F", "S", "N"} and code not in erlaubte_dienst_codes(person):
                    message = f"{name}, {column}: {code} ist für diese Person nicht erlaubt."
                elif code == "N" and str(person.get("nachtdienst", "Nein")) != "Ja":
                    message = f"{name}, {column}: Nachtdienst ist nicht aktiviert."
                else:
                    message = ""

            if message and message not in seen_errors:
                errors.append(message)
                seen_errors.add(message)
    return errors


def plan_to_paint_grid(plan, employees, settings) -> tuple[list[dict], list[dict]]:
    """Bereitet einen kompakten Monatsplan für den Pinsel-Editor vor."""
    normalised = normalise_plan(plan)
    stats = calculate_hours(normalised, employees, settings)
    stats_by_name = {
        str(row.get("Mitarbeitende", "")).strip().casefold(): row
        for row in stats.to_dict("records")
    }
    days = []
    plan_by_day = {}
    for entry in normalised:
        parsed = parse_date(entry.get("Datum"))
        if not parsed:
            continue
        key = day_column(parsed)
        plan_by_day[key] = entry
        days.append(
            {
                "key": key,
                "day": f"{parsed.day:02d}",
                "weekday": WOCHENTAGE[parsed.weekday()],
                "weekend": parsed.weekday() >= 5,
                "full_label": parsed.strftime("%d.%m.%Y"),
            }
        )

    protected_codes = {"U", "W", "FB", "K"}
    rows = []
    for person in employees:
        name = str(person.get("name", "")).strip()
        person_stats = stats_by_name.get(name.casefold(), {})
        target = float(person_stats.get("Soll", monthly_target(person)))
        actual = float(person_stats.get("Ist", 0.0))
        cells = {}
        locked = {}
        for item in days:
            code = code_for_person_day(plan_by_day.get(item["key"], {}), name)
            cells[item["key"]] = code
            locked[item["key"]] = code.upper() in protected_codes
        rows.append(
            {
                "name": name,
                "target": round(target, 1),
                "actual": round(actual, 1),
                "rest": round(target - actual, 1),
                "cells": cells,
                "locked": locked,
            }
        )
    return days, rows


def paint_result_to_matrix(result: dict, current_plan, employees, settings) -> pd.DataFrame:
    """Überträgt den Browser-Pinselstand in die bestehende Planmatrix."""
    matrix = plan_to_matrix(current_plan, employees, settings)
    result_rows = result.get("rows", []) if isinstance(result, dict) else []
    by_name = {
        str(row.get("name", "")).strip().casefold(): row
        for row in result_rows
        if isinstance(row, dict)
    }
    day_columns = [
        column for column in matrix.columns
        if re.match(r"^\d{2} ", str(column))
    ]
    for index, matrix_row in matrix.iterrows():
        name_key = str(matrix_row.get("Mitarbeitende", "")).strip().casefold()
        painted = by_name.get(name_key, {})
        cells = painted.get("cells", {}) if isinstance(painted, dict) else {}
        if not isinstance(cells, dict):
            continue
        for column in day_columns:
            if column not in cells:
                continue
            raw_code = str(cells.get(column, "") or "").strip()
            matrix.at[index, column] = "Frei" if raw_code.upper() == "FREI" else raw_code.upper()
    return matrix


def reset_month_editor_draft() -> None:
    """Verwirft eine alte Editor-Vorschau und erzwingt einen frischen Editor."""
    st.session_state.pop("dienstplan_editor_draft", None)
    st.session_state.pop("dienstplan_editor_draft_plan", None)
    st.session_state.pop("paint_grid_last_token", None)
    st.session_state.dienstplan_editor_version += 1



def replace_name_in_plan(plan, old_name: str, new_name: str):
    target = old_name.casefold()
    updated = normalise_plan(plan)
    for entry in updated:
        for column in CODE_TO_COLUMN.values():
            entry[column] = join_names(
                new_name if value.casefold() == target else value
                for value in split_names(entry.get(column, ""))
            )
    return updated

def remove_name_from_plan(plan, name: str):
    target = name.casefold()
    updated = normalise_plan(plan)
    for entry in updated:
        for column in CODE_TO_COLUMN.values():
            entry[column] = join_names(
                value
                for value in split_names(entry.get(column, ""))
                if value.casefold() != target
            )
    return updated

def shift_interval(day_value: date, code: str, settings: dict):
    prefixes = {"F": "f", "S": "s", "N": "n"}
    prefix = prefixes[code]
    start_dt = datetime.combine(day_value, parse_clock(settings[f"{prefix}_start"]))
    end_dt = datetime.combine(day_value, parse_clock(settings[f"{prefix}_ende"]))
    if end_dt <= start_dt:
        end_dt += timedelta(days=1)
    return start_dt, end_dt


def _validate_plan_impl(plan, employees, settings):
    errors = []
    warnings = []
    normalised = normalise_plan(plan)
    employee_map = {p["name"].casefold(): p for p in employees}
    employee_by_id = {int(person["id"]): person for person in employees}
    allowed_codes = {
        key: erlaubte_dienst_codes(person)
        for key, person in employee_map.items()
    }
    fachkraft_keys = {
        key for key, person in employee_map.items() if is_fachkraft(person)
    }
    assignments = {key: [] for key in employee_map}
    work_dates = {key: set() for key in employee_map}
    night_dates = {key: set() for key in employee_map}
    service_dates = {
        key: {"F": set(), "S": set(), "N": set()}
        for key in employee_map
    }
    service_counts = {
        key: {"F": 0, "S": 0, "N": 0}
        for key in employee_map
    }
    credited_absence_days = {key: 0 for key in employee_map}
    indexed_entries = []

    plan_columns = [
        ("Frühdienst", "F"), ("Spätdienst", "S"), ("Nachtdienst", "N"),
        ("Urlaub", "U"), ("Wunschfrei", "W"), ("Fortbildung", "FB"),
        ("Krank", "K"), ("Frei", "Frei"),
    ]

    for entry in normalised:
        day_value = parse_date(entry.get("Datum"))
        if not day_value:
            continue

        appearances = {}
        names_by_code = {
            code: split_names(entry.get(column, ""))
            for column, code in plan_columns
        }
        name_sets = {
            code: {name.casefold() for name in names}
            for code, names in names_by_code.items()
        }
        indexed_entries.append((day_value, name_sets))

        for _, code in plan_columns:
            for name in names_by_code[code]:
                key = name.casefold()
                appearances.setdefault(key, []).append(code)
                if key not in employee_map:
                    errors.append(f"Unbekannte Person im Plan: {name}")
                    continue
                if code in {"F", "S", "N"}:
                    person = employee_map[key]
                    if code not in allowed_codes[key]:
                        errors.append(
                            f"{name}: {code} passt nicht zur festen Dienstart "
                            f"„{feste_dienstart_text(person)}“ ({day_value.strftime('%d.%m.')})."
                        )
                    if code == "N" and str(person.get("nachtdienst", "Nein")) != "Ja":
                        errors.append(f"{name}: Nachtdienst ist nicht erlaubt ({day_value.strftime('%d.%m.')}).")
                    start_dt, end_dt = shift_interval(day_value, code, settings)
                    assignments[key].append((start_dt, end_dt, code))
                    work_dates[key].add(day_value)
                    service_dates[key][code].add(day_value)
                    if code == "N":
                        night_dates[key].add(day_value)

        for key, codes in appearances.items():
            if len(codes) > 1 and key in employee_map:
                errors.append(
                    f'{employee_map[key]["name"]}: Mehrere Einträge am {day_value.strftime("%d.%m.")}: '
                    + ", ".join(codes)
                )
            if key in employee_map and codes:
                first_code = codes[0]
                if first_code in {"F", "S", "N"}:
                    service_counts[key][first_code] += 1
                elif first_code in {"U", "FB", "K"} and day_value.weekday() < 5:
                    credited_absence_days[key] += 1

        actual = {
            code: len(names_by_code[code]) for code in ("F", "S", "N")
        }
        actual_fach = {
            code: len(name_sets[code] & fachkraft_keys)
            for code in ("F", "S", "N")
        }

        for code, prefix in [("F", "f"), ("S", "s"), ("N", "n")]:
            required = int(settings.get(f"min_{prefix}", 0))
            required_fach = int(settings.get(f"min_fach_{prefix}", 0))
            if actual[code] < required:
                warnings.append(f"{day_value.strftime('%d.%m.')}: {code} nur {actual[code]}/{required} besetzt.")
            if regel_aktiv(settings, "regel_fachkraft_aktiv") and actual_fach[code] < required_fach:
                errors.append(
                    f"{day_value.strftime('%d.%m.')}: {code} nur {actual_fach[code]}/{required_fach} Fachkräfte."
                )

    if regel_aktiv(settings, "regel_ruhezeit_aktiv"):
        min_rest = float(settings.get("ruhezeit", 11.0))
        for key, items in assignments.items():
            items.sort(key=lambda value: value[0])
            name = employee_map[key]["name"]
            for previous, current in zip(items, items[1:]):
                rest = (current[0] - previous[1]).total_seconds() / 3600.0
                if rest < min_rest:
                    errors.append(
                        f"{name}: Nur {rest:.1f} Stunden Ruhezeit zwischen {previous[2]} und {current[2]}."
                    )

    def longest_run(values: set[date]) -> int:
        longest = current = 0
        previous = None
        for value in sorted(values):
            current = current + 1 if previous and value == previous + timedelta(days=1) else 1
            longest = max(longest, current)
            previous = value
        return longest

    max_work = int(settings.get("max_arbeitstage", 6))
    max_nights = int(settings.get("max_nachtdienste", 4))
    if regel_aktiv(settings, "regel_max_arbeitstage_aktiv"):
        for key, dates in work_dates.items():
            if longest_run(dates) > max_work:
                errors.append(f'{employee_map[key]["name"]}: Mehr als {max_work} Arbeitstage hintereinander.')
    if regel_aktiv(settings, "regel_max_nachtdienste_aktiv"):
        for key, dates in night_dates.items():
            if longest_run(dates) > max_nights:
                errors.append(f'{employee_map[key]["name"]}: Mehr als {max_nights} Nachtdienste hintereinander.')

    max_weekends = max(0, int(settings.get("max_wochenenden", 2)))
    no_consecutive_weekends = bool(settings.get("keine_folgewochenenden", True))
    for key, dates in work_dates.items():
        weekends = worked_weekends(dates)
        if regel_aktiv(settings, "regel_max_wochenenden_aktiv") and len(weekends) > max_weekends:
            errors.append(
                f'{employee_map[key]["name"]}: An {len(weekends)} Wochenenden gearbeitet; '
                f'erlaubt sind höchstens {max_weekends}.'
            )
        if no_consecutive_weekends and has_consecutive_weekends(weekends):
            errors.append(
                f'{employee_map[key]["name"]}: Zwei direkt aufeinanderfolgende Wochenenden eingeplant.'
            )

    if regel_aktiv(settings, "regel_sollstunden_aktiv"):
        workdays = workdays_in_month(normalised)
        hours_by_code = {
            "F": float(settings.get("f_stunden", 0.0)),
            "S": float(settings.get("s_stunden", 0.0)),
            "N": float(settings.get("n_stunden", 0.0)),
        }
        for key, person in employee_map.items():
            target = monthly_target(person)
            absence_day_hours = target / workdays if workdays else 0.0
            actual = round(
                sum(service_counts[key][code] * hours_by_code[code] for code in ("F", "S", "N"))
                + credited_absence_days[key] * absence_day_hours,
                1,
            )
            if actual > target + 0.05:
                errors.append(
                    f'{person["name"]}: {actual:.1f} Iststunden überschreiten '
                    f'das Soll von {target:.1f} Stunden.'
                )

    for rule in eigene_planungsregeln(settings):
        if not bool(rule.get("aktiv", True)):
            continue
        rule_name = str(rule.get("name", "Eigene Regel"))
        rule_type = str(rule.get("typ", ""))

        if rule_type == "nicht_gemeinsam":
            first = employee_by_id.get(regel_zahl(rule, "mitarbeitende_a_id", -1))
            second = employee_by_id.get(regel_zahl(rule, "mitarbeitende_b_id", -1))
            rule_shift = str(rule.get("dienst", "ALL"))
            if not first or not second:
                continue
            for day_value, name_sets in indexed_entries:
                codes = [rule_shift] if rule_shift in {"F", "S", "N"} else ["F", "S", "N"]
                for code in codes:
                    names = name_sets[code]
                    if first["name"].casefold() in names and second["name"].casefold() in names:
                        errors.append(
                            f'{rule_name}: {first["name"]} und {second["name"]} sind '
                            f'am {day_value.strftime("%d.%m.") if day_value else "unbekannten Tag"} '
                            f'gemeinsam in {code}.'
                        )

        elif rule_type == "frei_nach_dienst":
            trigger_code = str(rule.get("dienst", "N"))
            free_days = max(1, regel_zahl(rule, "freie_tage", 1))
            for person in employees:
                if not regel_gilt_fuer_person(rule, person):
                    continue
                key = person["name"].casefold()
                trigger_days = service_dates[key].get(trigger_code, set())
                for trigger_day in trigger_days:
                    blocked_days = {
                        trigger_day + timedelta(days=offset)
                        for offset in range(1, free_days + 1)
                    }
                    if blocked_days & work_dates[key]:
                        errors.append(
                            f'{rule_name}: {person["name"]} benötigt nach {trigger_code} '
                            f'am {trigger_day.strftime("%d.%m.")} {free_days} freie Tage.'
                        )

        elif rule_type == "max_dienstfolge":
            target_code = str(rule.get("dienst", "N"))
            maximum = max(1, regel_zahl(rule, "maximum", 1))
            for person in employees:
                if not regel_gilt_fuer_person(rule, person):
                    continue
                key = person["name"].casefold()
                dates = service_dates[key].get(target_code, set())
                if longest_run(dates) > maximum:
                    errors.append(
                        f'{rule_name}: {person["name"]} hat mehr als {maximum} '
                        f'{target_code}-Dienste hintereinander.'
                    )

        elif rule_type == "nicht_an_wochentagen":
            blocked_weekdays = regel_wochentage(rule)
            if not blocked_weekdays:
                continue
            for person in employees:
                if not regel_gilt_fuer_person(rule, person):
                    continue
                key = person["name"].casefold()
                for work_day in sorted(work_dates[key]):
                    if work_day.weekday() in blocked_weekdays:
                        errors.append(
                            f'{rule_name}: {person["name"]} darf am '
                            f'{WOCHENTAGE[work_day.weekday()]}, {work_day.strftime("%d.%m.")}, '
                            "nicht arbeiten."
                        )

    errors = list(dict.fromkeys(errors))
    warnings = list(dict.fromkeys(warnings))
    return errors, warnings

def validate_plan(plan, employees, settings):
    return measured_call("Planprüfung", _validate_plan_impl, plan, employees, settings)


def _create_automatic_month_impl(year: int, month: int, employees, settings, vorgaben):
    """Bedarfsorientierte Planung wie in professionellen Personaleinsatzsystemen.

    1. Urlaub, Wunschfrei, Fortbildung und Krank werden zuerst als Abwesenheiten gesperrt.
    2. Der komplette Nachtdienst des Monats wird zuerst geplant.
    3. Danach werden die Wochenenden und erst anschließend die Werktage besetzt.
    4. Pro Person gelten höchstens zwei Arbeitswochenenden; Folgewochenenden sind gesperrt.
    5. Sollstunden und Dienstverteilung dienen zur fairen Auswahl.
    """
    plan = apply_vorgaben(year, month, employees, vorgaben)
    if not employees:
        return plan, {"assigned": 0, "fixed": 0, "shortages": ["Keine Mitarbeitenden vorhanden."]}

    employee_map = {str(person["name"]).casefold(): person for person in employees}
    employee_by_id = {int(person["id"]): person for person in employees}
    employee_order = {str(person["name"]).casefold(): index for index, person in enumerate(employees)}
    targets = {key: monthly_target(person) for key, person in employee_map.items()}
    allowed_codes = {
        key: erlaubte_dienst_codes(person)
        for key, person in employee_map.items()
    }
    planned_hours = {key: 0.0 for key in employee_map}
    assignments = {key: [] for key in employee_map}
    work_date_sets = {key: set() for key in employee_map}
    night_date_sets = {key: set() for key in employee_map}
    service_dates = {
        key: {"F": set(), "S": set(), "N": set()}
        for key in employee_map
    }
    service_counts = {key: {"F": 0, "S": 0, "N": 0} for key in employee_map}
    weekend_sets = {key: set() for key in employee_map}
    weekend_service_codes = {key: {} for key in employee_map}
    plan_dates = {id(entry): parse_date(entry.get("Datum")) for entry in plan}
    entry_codes = {}
    for entry in plan:
        codes = {}
        for code, column in CODE_TO_COLUMN.items():
            for name in split_names(entry.get(column, "")):
                codes.setdefault(name.casefold(), code)
        entry_codes[id(entry)] = codes

    # Bezahlte Abwesenheiten zählen bereits zum Monatssoll und müssen vor der
    # Dienstverteilung berücksichtigt werden. Wochenenden erhalten dabei keine
    # pauschalen Urlaubs-/Fortbildungs-/Krankstunden.
    workdays = workdays_in_month(plan)
    for entry in plan:
        day_value = plan_dates[id(entry)]
        if not day_value or day_value.weekday() >= 5:
            continue
        for key, code in entry_codes[id(entry)].items():
            if key in employee_map and code in {"U", "FB", "K"}:
                absence_day_hours = targets[key] / workdays if workdays else 0.0
                planned_hours[key] += absence_day_hours
    shortages = []
    assigned_total = 0
    fixed_total = 0

    code_to_hours = {
        "F": float(settings.get("f_stunden", 0.0)),
        "S": float(settings.get("s_stunden", 0.0)),
        "N": float(settings.get("n_stunden", 0.0)),
    }
    demand = {
        "F": max(0, int(settings.get("min_f", 0))),
        "S": max(0, int(settings.get("min_s", 0))),
        "N": max(0, int(settings.get("min_n", 0))),
    }
    fachkraft_rule_active = regel_aktiv(settings, "regel_fachkraft_aktiv")
    target_rule_active = regel_aktiv(settings, "regel_sollstunden_aktiv")
    rest_rule_active = regel_aktiv(settings, "regel_ruhezeit_aktiv")
    max_work_rule_active = regel_aktiv(settings, "regel_max_arbeitstage_aktiv")
    max_night_rule_active = regel_aktiv(settings, "regel_max_nachtdienste_aktiv")
    max_weekend_rule_active = regel_aktiv(settings, "regel_max_wochenenden_aktiv")
    night_block_rule_active = regel_aktiv(settings, "regel_nachtblock_aktiv")
    same_weekend_rule_active = regel_aktiv(settings, "regel_wochenende_gleich_aktiv")

    if fachkraft_rule_active:
        fach_demand = {
            "F": min(demand["F"], max(0, int(settings.get("min_fach_f", 0)))),
            "S": min(demand["S"], max(0, int(settings.get("min_fach_s", 0)))),
            "N": min(demand["N"], max(0, int(settings.get("min_fach_n", 0)))),
        }
    else:
        fach_demand = {"F": 0, "S": 0, "N": 0}
    min_rest = float(settings.get("ruhezeit", 11.0))
    max_work = int(settings.get("max_arbeitstage", 6))
    max_nights = int(settings.get("max_nachtdienste", 4))
    max_weekends = max(0, int(settings.get("max_wochenenden", 2)))
    no_consecutive_weekends = bool(settings.get("keine_folgewochenenden", True))
    active_custom_rules = [
        rule
        for rule in eigene_planungsregeln(settings)
        if bool(rule.get("aktiv", True))
    ]
    days_in_month = len(plan)

    static_candidates = {
        code: [
            person
            for person in employees
            if targets[str(person["name"]).casefold()] > 0
            and code_to_hours[code] > 0
            and code in allowed_codes[str(person["name"]).casefold()]
            and (code != "N" or str(person.get("nachtdienst", "Nein")) == "Ja")
        ]
        for code in ("F", "S", "N")
    }
    static_fach_candidates = {
        code: [person for person in static_candidates[code] if is_fachkraft(person)]
        for code in ("F", "S", "N")
    }

    def longest_run(values: set[date]) -> int:
        longest = current = 0
        previous = None
        for value in sorted(values):
            current = current + 1 if previous and value == previous + timedelta(days=1) else 1
            longest = max(longest, current)
            previous = value
        return longest

    def would_exceed_run(values: set[date], day_value: date, maximum: int) -> bool:
        if day_value in values:
            return longest_run(values) > maximum
        before = 0
        cursor = day_value - timedelta(days=1)
        while cursor in values:
            before += 1
            cursor -= timedelta(days=1)
        after = 0
        cursor = day_value + timedelta(days=1)
        while cursor in values:
            after += 1
            cursor += timedelta(days=1)
        return before + 1 + after > maximum

    def person_code(entry: dict, person: dict) -> str:
        return entry_codes[id(entry)].get(str(person["name"]).casefold(), "")

    def violates_custom_rule(person: dict, entry: dict, code: str) -> bool:
        key = str(person["name"]).casefold()
        day_value = plan_dates[id(entry)]
        if day_value is None:
            return True

        for rule in active_custom_rules:
            rule_type = str(rule.get("typ", ""))

            if rule_type == "nicht_gemeinsam":
                first_id = regel_zahl(rule, "mitarbeitende_a_id", -1)
                second_id = regel_zahl(rule, "mitarbeitende_b_id", -1)
                person_id = int(person.get("id", -1))
                if person_id not in {first_id, second_id}:
                    continue
                other_id = second_id if person_id == first_id else first_id
                other = employee_by_id.get(other_id)
                if other is None:
                    continue
                other_code = person_code(entry, other)
                rule_shift = str(rule.get("dienst", "ALL"))
                if other_code == code and (rule_shift == "ALL" or rule_shift == code):
                    return True

            elif rule_type == "frei_nach_dienst" and regel_gilt_fuer_person(rule, person):
                trigger_code = str(rule.get("dienst", "N"))
                free_days = max(1, regel_zahl(rule, "freie_tage", 1))
                for existing_day in service_dates[key].get(trigger_code, set()):
                    if existing_day < day_value <= existing_day + timedelta(days=free_days):
                        return True
                if code == trigger_code and any(
                    day_value < existing_day <= day_value + timedelta(days=free_days)
                    for existing_day in work_date_sets[key]
                ):
                    return True

            elif rule_type == "max_dienstfolge" and regel_gilt_fuer_person(rule, person):
                target_code = str(rule.get("dienst", "N"))
                maximum = max(1, regel_zahl(rule, "maximum", 1))
                if code == target_code and would_exceed_run(
                    service_dates[key].get(target_code, set()), day_value, maximum
                ):
                    return True

            elif rule_type == "nicht_an_wochentagen" and regel_gilt_fuer_person(rule, person):
                if day_value.weekday() in regel_wochentage(rule):
                    return True

        return False

    def register(person: dict, entry: dict, code: str, fixed: bool = False) -> None:
        nonlocal assigned_total, fixed_total
        key = str(person["name"]).casefold()
        day_value = plan_dates[id(entry)]
        if not day_value:
            return
        start_dt, end_dt = shift_interval(day_value, code, settings)
        assignments[key].append((day_value, start_dt, end_dt, code))
        work_date_sets[key].add(day_value)
        service_dates[key][code].add(day_value)
        if code == "N":
            night_date_sets[key].add(day_value)
        entry_codes[id(entry)][key] = code
        planned_hours[key] += code_to_hours[code]
        service_counts[key][code] += 1
        anchor = weekend_anchor(day_value)
        if anchor is not None:
            weekend_sets[key].add(anchor)
            weekend_service_codes[key].setdefault(anchor, set()).add(code)
        if fixed:
            fixed_total += 1
        else:
            assigned_total += 1

    # Bereits manuell im geöffneten Plan eingetragene Dienste registrieren.
    for entry in plan:
        for code, column in [("F", "Frühdienst"), ("S", "Spätdienst"), ("N", "Nachtdienst")]:
            for name in split_names(entry.get(column, "")):
                person = employee_map.get(name.casefold())
                if person:
                    register(person, entry, code, fixed=True)

    def can_assign(person: dict, entry: dict, code: str) -> bool:
        key = str(person["name"]).casefold()
        day_value = plan_dates[id(entry)]
        if day_value is None or targets[key] <= 0 or code_to_hours[code] <= 0:
            return False
        # Sollstunden sind eine feste Obergrenze. Die Automatik lässt einen
        # Dienst offen, statt die Person über ihr persönliches Monatssoll zu planen.
        if (
            target_rule_active
            and planned_hours[key] + code_to_hours[code] > targets[key] + 0.001
        ):
            return False
        if person_code(entry, person):
            return False
        # Eine feste Dienstart schränkt nur die erlaubte Schicht ein. Sie erzeugt
        # keine tägliche Einplanung und ändert weder Sollstunden noch freie Tage.
        if code not in allowed_codes[key]:
            return False
        if code == "N" and str(person.get("nachtdienst", "Nein")) != "Ja":
            return False

        candidate_start, candidate_end = shift_interval(day_value, code, settings)
        for _, existing_start, existing_end, _ in assignments[key]:
            if candidate_start >= existing_end:
                rest = (candidate_start - existing_end).total_seconds() / 3600.0
            elif existing_start >= candidate_end:
                rest = (existing_start - candidate_end).total_seconds() / 3600.0
            else:
                return False
            if rest_rule_active and rest < min_rest:
                return False

        if max_work_rule_active and would_exceed_run(work_date_sets[key], day_value, max_work):
            return False
        if code == "N" and max_night_rule_active:
            if would_exceed_run(night_date_sets[key], day_value, max_nights):
                return False

        anchor = weekend_anchor(day_value)
        if anchor is not None:
            candidate_weekends = set(weekend_sets[key]) | {anchor}
            if max_weekend_rule_active and len(candidate_weekends) > max_weekends:
                return False
            if no_consecutive_weekends and has_consecutive_weekends(candidate_weekends):
                return False
        if violates_custom_rule(person, entry, code):
            return False
        return True

    def assign(person: dict, entry: dict, code: str) -> None:
        column = CODE_TO_COLUMN[code]
        entry[column] = join_names(split_names(entry.get(column, "")) + [person["name"]])
        register(person, entry, code, fixed=False)

    night_capable_target = sum(
        targets[str(person["name"]).casefold()]
        for person in employees
        if str(person.get("nachtdienst", "Nein")) == "Ja"
        and targets[str(person["name"]).casefold()] > 0
    )
    total_night_slots = days_in_month * demand["N"]

    only_shift_labels = {
        "F": "Nur Frühdienst",
        "S": "Nur Spätdienst",
        "N": "Nur Nachtdienst",
    }

    def shift_priority(person: dict, code: str) -> int:
        """Mitarbeitende mit nur einer Dienstart werden zuerst eingeplant.

        Dadurch werden zum Beispiel reine Nachtwachen nicht von Mitarbeitenden
        mit „Alle Dienste“ aus den Nachtdiensten verdrängt. Erst wenn diese
        Personen nicht verfügbar sind, nutzt die Automatik flexibel einsetzbare
        Mitarbeitende.
        """
        return 0 if feste_dienstart_text(person) == only_shift_labels[code] else 1

    def weekend_pair_priority(person: dict, day_value: date, code: str) -> int:
        """Am Sonntag möglichst dieselben Personen und Schichten wie Samstag nutzen."""
        if not same_weekend_rule_active:
            return 0
        anchor = weekend_anchor(day_value)
        if anchor is None:
            return 0

        key = str(person["name"]).casefold()
        same_weekend_codes = weekend_service_codes[key].get(anchor, set())
        if code in same_weekend_codes:
            return 0
        if same_weekend_codes:
            return 1
        return 2

    def score_candidate(person: dict, entry: dict, code: str):
        key = str(person["name"]).casefold()
        day_value = plan_dates[id(entry)]
        target = max(targets[key], 0.01)
        projected = planned_hours[key] + code_to_hours[code]
        projected_utilisation = projected / target
        over_ratio = max(0.0, projected - target) / target
        previous_night = day_value - timedelta(days=1) in night_date_sets[key]
        next_night = day_value + timedelta(days=1) in night_date_sets[key]
        night_block_penalty = (
            0
            if (
                not night_block_rule_active
                or code != "N"
                or previous_night
                or next_night
            )
            else 1
        )
        expected_nights = (
            total_night_slots * target / night_capable_target
            if code == "N" and night_capable_target > 0
            else 1.0
        )
        night_load = (
            (service_counts[key]["N"] + 1) / max(expected_nights, 0.5)
            if code == "N"
            else 0.0
        )
        weekend_load = len(weekend_sets[key])
        weekend_pair = weekend_pair_priority(person, day_value, code)
        restricted_shift = shift_priority(person, code)
        rotation = (day_value.day - 1 - employee_order[key] * 2) % max(1, days_in_month)

        # Für Nachtdienste gilt eine feste Priorität:
        # 1. Mitarbeitende mit „Nur Nachtdienst"
        # 2. bereits begonnene Nachtblöcke fortsetzen
        # 3. innerhalb dieser Gruppen fair nach Sollstunden verteilen
        if code == "N":
            return (
                restricted_shift,
                night_block_penalty,
                round(projected_utilisation, 6),
                round(night_load, 6),
                round(over_ratio, 6),
                service_counts[key][code],
                sum(service_counts[key].values()),
                rotation,
                str(person["name"]).casefold(),
            )

        # Früh- und Spätdienste werden erst nach allen Nachtdiensten geplant.
        # Am Wochenende bleiben die Personen nach Möglichkeit an beiden Tagen
        # in derselben Schicht.
        return (
            restricted_shift,
            weekend_pair,
            round(projected_utilisation, 6),
            round(over_ratio, 6),
            weekend_load if day_value.weekday() >= 5 else 0,
            service_counts[key][code],
            sum(service_counts[key].values()),
            rotation,
            str(person["name"]).casefold(),
        )

    def fill_slots(entry: dict, code: str, needed: int, fach_only: bool) -> None:
        candidate_pool = static_fach_candidates[code] if fach_only else static_candidates[code]
        for _ in range(max(0, needed)):
            candidates = [
                person for person in candidate_pool
                if can_assign(person, entry, code)
            ]
            if not candidates:
                kind = "Fachkraft" if fach_only else "Person"
                shortages.append(f'{entry["Datum"]}: {code} – keine passende {kind} verfügbar.')
                return
            selected = min(candidates, key=lambda person: score_candidate(person, entry, code))
            assign(selected, entry, code)

    def fill_fach(entry: dict, code: str) -> None:
        current_names = split_names(entry.get(CODE_TO_COLUMN[code], ""))
        current_fach = sum(
            1 for name in current_names
            if name.casefold() in employee_map and is_fachkraft(employee_map[name.casefold()])
        )
        fill_slots(entry, code, fach_demand[code] - current_fach, fach_only=True)

    def fill_total(entry: dict, code: str) -> None:
        current_total = len(split_names(entry.get(CODE_TO_COLUMN[code], "")))
        fill_slots(entry, code, demand[code] - current_total, fach_only=False)

    weekend_entries = [
        entry for entry in plan
        if (day_value := plan_dates[id(entry)]) and day_value.weekday() >= 5
    ]
    weekday_entries = [
        entry for entry in plan
        if (day_value := plan_dates[id(entry)]) and day_value.weekday() < 5
    ]
    planning_days = (
        weekend_entries + weekday_entries
        if regel_aktiv(settings, "regel_wochenende_zuerst_aktiv")
        else list(plan)
    )

    if regel_aktiv(settings, "regel_nacht_zuerst_aktiv"):
        for entry in plan:
            fill_fach(entry, "N")
            fill_total(entry, "N")
        for entry in planning_days:
            for code in ("F", "S"):
                fill_fach(entry, code)
                fill_total(entry, code)
            entry["Status"] = "Nachtdienste zuerst geplant"
    else:
        for entry in planning_days:
            for code in ("F", "S", "N"):
                fill_fach(entry, code)
                fill_total(entry, code)
            entry["Status"] = "Nach aktiven Planungsregeln geplant"

    required_hours = days_in_month * sum(demand[c] * code_to_hours[c] for c in ("F", "S", "N"))
    available_hours = sum(targets.values())
    return plan, {
        "assigned": assigned_total,
        "fixed": fixed_total,
        "shortages": list(dict.fromkeys(shortages)),
        "required_hours": round(required_hours, 1),
        "available_hours": round(available_hours, 1),
    }


def create_automatic_month(year: int, month: int, employees, settings, vorgaben):
    return measured_call(
        "Automatische Planung",
        _create_automatic_month_impl,
        year,
        month,
        employees,
        settings,
        vorgaben,
    )


def safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9ÄÖÜäöüß_-]+", "_", str(name).strip())
    return cleaned.strip("_") or "Dienstplan"


def export_matrix(plan, employees, settings):
    return plan_to_matrix(plan, employees, settings)


def create_excel(plan, plan_name, employees, settings):
    matrix = export_matrix(plan, employees, settings)
    year, month = month_from_plan(plan)
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dienstplan"

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(matrix.columns))
    title = sheet.cell(1, 1, plan_name)
    title.font = Font(size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="203040")
    title.alignment = Alignment(horizontal="center")

    for column_index, column in enumerate(matrix.columns, start=1):
        cell = sheet.cell(3, column_index, column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(matrix.itertuples(index=False), start=4):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_index, column in enumerate(matrix.columns, start=1):
        width = 22 if column == "Mitarbeitende" else 7
        if re.match(r"^\d{2} ", str(column)):
            day_number = int(str(column)[:2])
            if date(year, month, day_number).weekday() >= 5:
                for row_index in range(3, len(matrix) + 4):
                    sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor="FFF2CC")
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def create_pdf(plan, plan_name, employees, settings):
    matrix = export_matrix(plan, employees, settings)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A3),
        leftMargin=7 * mm,
        rightMargin=7 * mm,
        topMargin=7 * mm,
        bottomMargin=7 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#172033"),
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=5.7,
        leading=6.4,
        alignment=1,
    )
    header_style = ParagraphStyle(
        "Header",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    data = [[Paragraph(escape(str(column)), header_style) for column in matrix.columns]]
    for _, row in matrix.iterrows():
        data.append([Paragraph(escape(str(value)), cell_style) for value in row])

    widths = [32 * mm if column == "Mitarbeitende" else 8.2 * mm for column in matrix.columns]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#44546A")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C8CED8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 1.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 1.5),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    document.build([Paragraph(escape(plan_name), title_style), table])
    output.seek(0)
    return output.getvalue()


# -----------------------------------------------------------------------------
# Neue Benutzeroberfläche
# -----------------------------------------------------------------------------

NAVIGATION = {
    "🏠  Übersicht": "Übersicht",
    "👥  Mitarbeitende": "Mitarbeitende",
    "📅  Abwesenheiten": "Abwesenheiten",
    "📐  Planungsregeln": "Planungsregeln",
    "🧩  Vorlagen": "Vorlagen",
    "🗓️  Dienstplan": "Dienstplan",
    "📊  Auswertung": "Auswertung",
    "💬  Kontakt": "Feedback & Hilfe",
    "⚖️  Rechtliches": "Rechtliches",
    "⚙️  Einstellungen": "Einstellungen",
}

BUILTIN_TEMPLATES = {
    "Standardstation": {
        "description": "Ausgewogene Grundbesetzung für eine normale Pflegestation.",
        "settings": {
            "min_f": 4,
            "min_s": 3,
            "min_n": 1,
            "min_fach_f": 1,
            "min_fach_s": 1,
            "min_fach_n": 1,
            "max_arbeitstage": 6,
            "max_nachtdienste": 3,
            "max_wochenenden": 2,
            "keine_folgewochenenden": True,
        },
    },
    "Hoher Pflegebedarf": {
        "description": "Mehr Personal im Früh- und Spätdienst sowie mindestens zwei Fachkräfte am Tag.",
        "settings": {
            "min_f": 5,
            "min_s": 3,
            "min_n": 1,
            "min_fach_f": 2,
            "min_fach_s": 2,
            "min_fach_n": 1,
            "max_arbeitstage": 6,
            "max_nachtdienste": 3,
            "max_wochenenden": 2,
            "keine_folgewochenenden": True,
        },
    },
    "Kleine Station": {
        "description": "Reduzierte Mindestbesetzung für kleinere Bereiche oder Übergangszeiten.",
        "settings": {
            "min_f": 3,
            "min_s": 2,
            "min_n": 1,
            "min_fach_f": 1,
            "min_fach_s": 1,
            "min_fach_n": 1,
            "max_arbeitstage": 6,
            "max_nachtdienste": 3,
            "max_wochenenden": 2,
            "keine_folgewochenenden": True,
        },
    },
}


def section_title(title: str, help_text: str = "") -> None:
    st.markdown(f'<div class="np-section-title">{escape(title)}</div>', unsafe_allow_html=True)
    if help_text:
        st.markdown(f'<div class="np-help">{escape(help_text)}</div>', unsafe_allow_html=True)


def kpi_card(label: str, value: str, note: str = "") -> None:
    st.markdown(
        f"""
        <div class="np-kpi">
            <div class="np-kpi-label">{escape(label)}</div>
            <div class="np-kpi-value">{escape(value)}</div>
            <div class="np-kpi-note">{escape(note)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def german_date(value: str) -> str:
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return str(value)


def german_datetime(value: str) -> str:
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.strftime("%d.%m.%Y · %H:%M Uhr")
    except (TypeError, ValueError):
        return str(value)


def render_employee_table(items: list[dict]) -> None:
    if not items:
        st.info("Noch keine Mitarbeitenden angelegt.")
        return

    rows = []
    for person in items:
        night_class = "np-badge-green" if str(person.get("nachtdienst")) == "Ja" else "np-badge-gray"
        night_text = "Erlaubt" if str(person.get("nachtdienst")) == "Ja" else "Nicht erlaubt"
        fixed_shift = feste_dienstart_text(person)
        shift_class = "np-badge-blue" if fixed_shift == "Alle Dienste" else "np-badge-orange"
        rows.append(
            "<tr>"
            f'<td><span class="np-name">{escape(str(person["name"]))}</span></td>'
            f'<td>{escape(str(person.get("qualifikation", "")))}</td>'
            f'<td>{escape(percentage_text(monthly_target(person)))}</td>'
            f'<td>{monthly_target(person):.1f} Std.</td>'
            f'<td><span class="np-badge {night_class}">{night_text}</span></td>'
            f'<td><span class="np-badge {shift_class}">{escape(fixed_shift)}</span></td>'
            "</tr>"
        )

    st.markdown(
        """
        <div class="np-table-wrap">
            <table class="np-table">
                <thead>
                    <tr>
                        <th>Name</th>
                        <th>Qualifikation</th>
                        <th>Beschäftigung</th>
                        <th>Sollstunden</th>
                        <th>Nachtdienst</th>
                        <th>Erlaubte Dienste</th>
                    </tr>
                </thead>
                <tbody>
        """
        + "".join(rows)
        + """
                </tbody>
            </table>
        </div>
        """,
        unsafe_allow_html=True,
    )


def readable_plan_matrix(matrix: pd.DataFrame):
    day_columns = [column for column in matrix.columns if re.match(r"^\d{2} ", str(column))]

    def style_code(value):
        code = str(value or "").strip().upper()
        styles = {
            "F": "background-color: #e8f3ff; color: #174ea6; font-weight: 700; text-align: center;",
            "S": "background-color: #fff2df; color: #9a4f00; font-weight: 700; text-align: center;",
            "N": "background-color: #ede9fe; color: #5b21b6; font-weight: 700; text-align: center;",
            "U": "background-color: #e8f8f2; color: #087a54; font-weight: 700; text-align: center;",
            "W": "background-color: #f1f5f9; color: #475467; font-weight: 700; text-align: center;",
            "FB": "background-color: #e0f2fe; color: #075985; font-weight: 700; text-align: center;",
            "K": "background-color: #feecec; color: #b42318; font-weight: 700; text-align: center;",
            "FREI": "background-color: #f4f4f5; color: #52525b; font-weight: 650; text-align: center;",
        }
        return styles.get(code, "text-align: center;")

    styled = matrix.style
    if day_columns:
        styled = styled.map(style_code, subset=day_columns)
    numeric = [column for column in ["Soll", "Ist", "Rest"] if column in matrix.columns]
    if numeric:
        styled = styled.format({column: "{:.1f}" for column in numeric})
    return styled


def settings_snapshot(values: dict) -> dict:
    keys = [
        "vollzeit_monatssoll",
        "f_start", "f_ende", "f_stunden", "min_f", "min_fach_f",
        "s_start", "s_ende", "s_stunden", "min_s", "min_fach_s",
        "n_start", "n_ende", "n_stunden", "min_n", "min_fach_n",
        "ruhezeit", "max_arbeitstage", "max_nachtdienste", "max_wochenenden",
        "keine_folgewochenenden", "max_plusstunden",
        "regel_ruhezeit_aktiv", "regel_max_arbeitstage_aktiv",
        "regel_max_nachtdienste_aktiv", "regel_max_wochenenden_aktiv",
        "regel_sollstunden_aktiv", "regel_fachkraft_aktiv",
        "regel_nacht_zuerst_aktiv", "regel_nachtblock_aktiv",
        "regel_wochenende_zuerst_aktiv", "regel_wochenende_gleich_aktiv",
        "eigene_planungsregeln",
    ]
    return {key: values.get(key) for key in keys if key in values}


def apply_template_settings(template_values: dict) -> None:
    new_settings = {**st.session_state.settings, **template_values}
    save_einstellungen(new_settings)
    st.session_state.settings = new_settings


def copy_plan_to_month(source_plan, year: int, month: int, employees, vorgaben):
    """Kopiert Dienste nach Kalendertag. Neue Abwesenheiten haben Vorrang."""
    target_plan = apply_vorgaben(year, month, employees, vorgaben)
    employee_names = {str(person["name"]).casefold(): str(person["name"]) for person in employees}
    source_by_day = {int(row.get("Tag", index)): row for index, row in enumerate(normalise_plan(source_plan), start=1)}

    for target_entry in target_plan:
        source_entry = source_by_day.get(int(target_entry["Tag"]))
        if not source_entry:
            continue
        for column in ["Frühdienst", "Spätdienst", "Nachtdienst", "Frei"]:
            for source_name in split_names(source_entry.get(column, "")):
                real_name = employee_names.get(source_name.casefold())
                if not real_name:
                    continue
                if code_for_person_day(target_entry, real_name):
                    continue
                target_entry[column] = join_names(split_names(target_entry.get(column, "")) + [real_name])
        target_entry["Status"] = "Aus Vorlage übernommen"
    return target_plan


def show_plan_legend() -> None:
    st.markdown(
        """
        <div style="display:flex;gap:.4rem;flex-wrap:wrap;margin:.25rem 0 1rem 0;">
            <span class="np-badge np-badge-blue">F · Frühdienst</span>
            <span class="np-badge np-badge-orange">S · Spätdienst</span>
            <span class="np-badge" style="background:#ede9fe;color:#5b21b6;">N · Nachtdienst</span>
            <span class="np-badge np-badge-green">U · Urlaub</span>
            <span class="np-badge np-badge-gray">W · Wunschfrei</span>
            <span class="np-badge" style="background:#e0f2fe;color:#075985;">FB · Fortbildung</span>
            <span class="np-badge np-badge-red">K · Krank</span>
            <span class="np-badge np-badge-gray">Frei · Arbeitsfrei</span>
        </div>
        """,
        unsafe_allow_html=True,
    )


st.sidebar.markdown(
    """
    <div class="np-brand">
        <div class="np-brand-name">🏥 NursePlan Pro</div>
        <div class="np-brand-subtitle">Dienstplanung für Pflegeteams</div>
    </div>
    """,
    unsafe_allow_html=True,
)
st.sidebar.caption(f"Angemeldet: {current_user_email()}")
if st.sidebar.button("Abmelden", width="stretch"):
    st.session_state.auth_logout_requested = True
    delete_auth_cookie()
    try:
        sign_out_user()
    finally:
        for state_key in list(st.session_state):
            if state_key != "auth_logout_requested":
                del st.session_state[state_key]
    st.rerun()

selected_navigation = st.sidebar.radio(
    "Navigation",
    list(NAVIGATION),
    label_visibility="collapsed",
    key="main_navigation",
)
page = NAVIGATION[selected_navigation]
st.sidebar.divider()
st.sidebar.caption("Tipp: Zuerst Mitarbeitende und Abwesenheiten pflegen. Danach den Dienstplan erstellen.")

settings = st.session_state.settings
employees, vorgaben, saved_plans, saved_templates = load_account_overview_data()


if page == "Übersicht":
    page_header(
        "Übersicht",
        "Alle wichtigen Informationen auf einen Blick. Von hier aus erkennst du, was vor der Planung noch fehlt.",
    )

    open_plan = st.session_state.dienstplan_vorlage
    open_plan_name = st.session_state.dienstplan_name if open_plan else "Kein Plan geöffnet"
    active_absences = len(vorgaben)
    night_capable = sum(1 for person in employees if str(person.get("nachtdienst")) == "Ja")

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        kpi_card("Mitarbeitende", str(len(employees)), "im Team gespeichert")
    with c2:
        kpi_card("Nachtdienst möglich", str(night_capable), "Mitarbeitende")
    with c3:
        kpi_card("Abwesenheiten", str(active_absences), "gespeicherte Zeiträume")
    with c4:
        kpi_card("Gespeicherte Pläne", str(len(saved_plans)), "jederzeit wieder ladbar")

    left, right = st.columns([1.35, 1], gap="large")
    with left:
        section_title("Aktueller Arbeitsstand", "Hier siehst du, ob bereits ein Monatsplan geöffnet ist.")
        if open_plan:
            plan_year, plan_month = month_from_plan(open_plan)
            errors, warnings = validate_plan(open_plan, employees, settings)
            st.markdown(
                f"""
                <div class="np-card">
                    <div style="font-size:.8rem;color:#667085;font-weight:700;">GEÖFFNETER PLAN</div>
                    <div style="font-size:1.25rem;font-weight:800;margin:.2rem 0;">{escape(open_plan_name)}</div>
                    <div style="color:#667085;">{MONATSNAMEN[plan_month]} {plan_year} · {len(errors)} Fehler · {len(warnings)} Hinweise</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        else:
            st.info("Noch kein Dienstplan geöffnet. Öffne die Seite „Dienstplan“, um einen Monat zu erstellen.")

        section_title("Vor der automatischen Planung prüfen")
        checklist = [
            (bool(employees), "Mitarbeitende sind angelegt"),
            (night_capable > 0, "Mindestens eine Person darf Nachtdienst übernehmen"),
            (settings.get("min_f", 0) > 0, "Mindestbesetzung ist eingestellt"),
            (True, "Urlaub und Wünsche wurden geprüft"),
        ]
        for completed, text in checklist:
            icon = "✅" if completed else "⚠️"
            st.write(f"{icon} {text}")

    with right:
        section_title("Aktuelle Besetzung", "Diese Werte verwendet die automatische Planung.")
        st.markdown(
            f"""
            <div class="np-card">
                <div style="display:grid;grid-template-columns:1fr auto;gap:.7rem;">
                    <span>Frühdienst</span><strong>{int(settings.get('min_f', 0))} Personen</strong>
                    <span>Spätdienst</span><strong>{int(settings.get('min_s', 0))} Personen</strong>
                    <span>Nachtdienst</span><strong>{int(settings.get('min_n', 0))} Person(en)</strong>
                    <span>Max. Arbeitstage</span><strong>{int(settings.get('max_arbeitstage', 0))}</strong>
                    <span>Max. Wochenenden</span><strong>{int(settings.get('max_wochenenden', 0))}</strong>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        section_title("Nächste Schritte")
        st.markdown(
            """
            <div class="np-card">
                <strong>1.</strong> Teamdaten prüfen<br>
                <strong>2.</strong> Urlaub und Frei-Wünsche eintragen<br>
                <strong>3.</strong> Vorlage oder Besetzung auswählen<br>
                <strong>4.</strong> Dienstplan erstellen und prüfen
            </div>
            """,
            unsafe_allow_html=True,
        )

    if employees:
        section_title("Team nach Qualifikation")
        qualification_counts = (
            pd.DataFrame(employees)["qualifikation"]
            .fillna("Nicht angegeben")
            .value_counts()
            .rename_axis("Qualifikation")
            .reset_index(name="Anzahl")
        )
        st.bar_chart(qualification_counts.set_index("Qualifikation"))


elif page == "Mitarbeitende":
    page_header(
        "Mitarbeitende",
        "Lege das Team einmal sauber an. Sollstunden, Qualifikation und erlaubte Dienste werden später automatisch berücksichtigt.",
    )

    total_target = sum(monthly_target(person) for person in employees)
    specialists = sum(1 for person in employees if is_fachkraft(person))
    c1, c2, c3 = st.columns(3)
    with c1:
        kpi_card("Teamgröße", str(len(employees)), "aktive Mitarbeitende")
    with c2:
        kpi_card("Fachkräfte", str(specialists), "für die Mindestbesetzung")
    with c3:
        kpi_card("Monatssoll gesamt", f"{total_target:.1f}", "Stunden im Team")

    add_tab, edit_tab, overview_tab = st.tabs(["➕ Neue Person", "✏️ Bearbeiten", "👥 Teamübersicht"])

    with add_tab:
        section_title("Neue Person anlegen", "Trage nur die Angaben ein, die für die Dienstplanung wirklich benötigt werden.")
        with st.form("new_employee_form", clear_on_submit=True):
            left, right = st.columns(2, gap="large")
            with left:
                new_name = st.text_input("Name", placeholder="Zum Beispiel: Anna Müller")
                new_target = st.number_input(
                    "Sollstunden im Monat",
                    min_value=0.0,
                    max_value=300.0,
                    value=float(settings.get("vollzeit_monatssoll", 167.0)),
                    step=0.5,
                    help="Die Stunden laut Arbeitsvertrag für den ausgewählten Monat.",
                )
                st.markdown(
                    f"""
                    <div class="np-percent">
                        <div class="np-percent-label">Berechneter Beschäftigungsumfang</div>
                        <div class="np-percent-value">{percentage_text(new_target)}</div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
            with right:
                new_qualification = st.selectbox(
                    "Qualifikation",
                    QUALIFIKATIONEN,
                    help="Die Qualifikation wird für die Fachkraftbesetzung geprüft.",
                )
                new_fixed_shift = st.selectbox(
                    "Welche Dienste sind erlaubt?",
                    FESTE_DIENSTARTEN,
                    help="Wähle „Alle Dienste“, wenn Früh-, Spät- und Nachtdienst möglich sind.",
                )
                new_night = st.checkbox(
                    "Nachtdienst ist erlaubt",
                    value=True,
                    help="Deaktivieren, wenn die Person grundsätzlich keine Nachtdienste übernehmen darf.",
                )
            submitted = st.form_submit_button("Mitarbeitende speichern", type="primary", width="stretch")

        if submitted:
            if not new_name.strip():
                st.warning("Bitte trage einen Namen ein.")
            elif new_target <= 0:
                st.warning("Die Sollstunden müssen größer als 0 sein.")
            elif new_fixed_shift == "Nur Nachtdienst" and not new_night:
                st.warning("Bei „Nur Nachtdienst“ muss Nachtdienst erlaubt sein.")
            elif any(person["name"].strip().casefold() == new_name.strip().casefold() for person in employees):
                st.warning("Dieser Name ist bereits vorhanden.")
            else:
                add_mitarbeitende(
                    name=new_name.strip(),
                    stellenanteil=stored_percentage(new_target),
                    wochenstunden=weekly_from_month(new_target),
                    qualifikation=new_qualification,
                    erfahrung="Nicht erfasst",
                    nachtdienst="Ja" if new_night else "Nein",
                    monatssollstunden=float(new_target),
                    schutzstatus="Keiner",
                    feste_dienstart=new_fixed_shift,
                )
                invalidate_account_cache("Mitarbeitende")
                st.success(f"{new_name.strip()} wurde gespeichert.")
                st.rerun()

    with edit_tab:
        if not employees:
            st.info("Lege zuerst eine Person an.")
        else:
            labels = {person["name"]: person for person in employees}
            selected_name = st.selectbox("Mitarbeitende auswählen", list(labels), key="employee_to_edit")
            selected = labels[selected_name]

            with st.form(f'edit_employee_form_{selected["id"]}'):
                left, right = st.columns(2, gap="large")
                with left:
                    edit_name = st.text_input("Name", value=selected["name"])
                    edit_target = st.number_input(
                        "Sollstunden im Monat",
                        min_value=0.0,
                        max_value=300.0,
                        value=float(monthly_target(selected)),
                        step=0.5,
                    )
                    st.markdown(
                        f"""
                        <div class="np-percent">
                            <div class="np-percent-label">Berechneter Beschäftigungsumfang</div>
                            <div class="np-percent-value">{percentage_text(edit_target)}</div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                with right:
                    qualification_index = (
                        QUALIFIKATIONEN.index(selected["qualifikation"])
                        if selected.get("qualifikation") in QUALIFIKATIONEN
                        else 0
                    )
                    edit_qualification = st.selectbox("Qualifikation", QUALIFIKATIONEN, index=qualification_index)
                    current_fixed_shift = feste_dienstart_text(selected)
                    edit_fixed_shift = st.selectbox(
                        "Welche Dienste sind erlaubt?",
                        FESTE_DIENSTARTEN,
                        index=FESTE_DIENSTARTEN.index(current_fixed_shift),
                    )
                    edit_night = st.checkbox(
                        "Nachtdienst ist erlaubt",
                        value=str(selected.get("nachtdienst", "Nein")) == "Ja",
                    )
                save_employee = st.form_submit_button("Änderungen speichern", type="primary", width="stretch")

            if save_employee:
                if not edit_name.strip():
                    st.warning("Der Name darf nicht leer sein.")
                elif edit_target <= 0:
                    st.warning("Die Sollstunden müssen größer als 0 sein.")
                elif edit_fixed_shift == "Nur Nachtdienst" and not edit_night:
                    st.warning("Bei „Nur Nachtdienst“ muss Nachtdienst erlaubt sein.")
                elif any(
                    person["id"] != selected["id"]
                    and person["name"].strip().casefold() == edit_name.strip().casefold()
                    for person in employees
                ):
                    st.warning("Dieser Name ist bereits vorhanden.")
                else:
                    update_mitarbeitende(
                        mitarbeitende_id=int(selected["id"]),
                        name=edit_name.strip(),
                        stellenanteil=stored_percentage(edit_target),
                        wochenstunden=weekly_from_month(edit_target),
                        qualifikation=edit_qualification,
                        erfahrung=str(selected.get("erfahrung", "Nicht erfasst")),
                        nachtdienst="Ja" if edit_night else "Nein",
                        monatssollstunden=float(edit_target),
                        schutzstatus=str(selected.get("schutzstatus", "Keiner")),
                        feste_dienstart=edit_fixed_shift,
                    )
                    invalidate_account_cache("Mitarbeitende")
                    if selected["name"] != edit_name.strip():
                        st.session_state.dienstplan_vorlage = replace_name_in_plan(
                            st.session_state.dienstplan_vorlage,
                            selected["name"],
                            edit_name.strip(),
                        )
                    st.success("Die Änderungen wurden gespeichert.")
                    st.rerun()

            st.divider()
            st.markdown("**Person löschen**")
            confirm_delete = st.checkbox(
                f'Ich möchte {selected["name"]} wirklich löschen.',
                key=f'confirm_delete_{selected["id"]}',
            )
            if st.button("Person löschen", disabled=not confirm_delete, width="stretch"):
                delete_mitarbeitende(int(selected["id"]))
                invalidate_account_cache("Mitarbeitende", "Abwesenheiten")
                st.session_state.dienstplan_vorlage = remove_name_from_plan(
                    st.session_state.dienstplan_vorlage,
                    selected["name"],
                )
                st.rerun()

    with overview_tab:
        filter_left, filter_right = st.columns([1.3, 1])
        with filter_left:
            search_name = st.text_input("Suchen", placeholder="Name eingeben", key="employee_search")
        with filter_right:
            qualification_filter = st.selectbox(
                "Qualifikation filtern",
                ["Alle"] + QUALIFIKATIONEN,
                key="employee_qualification_filter",
            )
        filtered = [
            person for person in employees
            if (not search_name.strip() or search_name.strip().casefold() in person["name"].casefold())
            and (qualification_filter == "Alle" or person.get("qualifikation") == qualification_filter)
        ]
        render_employee_table(filtered)
        st.caption(f"Angezeigt: {len(filtered)} von {len(employees)} Mitarbeitenden")


elif page == "Abwesenheiten":
    page_header(
        "Abwesenheiten und Wünsche",
        "Trage Urlaub, Wunschfrei, Fortbildung oder Krankheit ein. Diese Zeiten werden bei der automatischen Planung geschützt.",
    )

    if not employees:
        st.warning("Bitte lege zuerst Mitarbeitende an.")
    else:
        add_tab, overview_tab = st.tabs(["➕ Eintragen", "📋 Übersicht"])
        with add_tab:
            section_title("Neue Abwesenheit", "Wähle eine Person, die Art und den genauen Zeitraum.")
            with st.form("absence_form", clear_on_submit=True):
                c1, c2 = st.columns(2)
                employee_options = {person["name"]: int(person["id"]) for person in employees}
                label_to_code = {label: code for code, label in VORGABE_LABELS.items()}
                with c1:
                    selected_employee = st.selectbox("Mitarbeitende", list(employee_options))
                    selected_label = st.selectbox("Art der Abwesenheit", list(label_to_code))
                with c2:
                    start_value = st.date_input("Von", value=date.today(), key="absence_start")
                    end_value = st.date_input("Bis einschließlich", value=date.today(), key="absence_end")
                note = st.text_input("Bemerkung", placeholder="Optional, zum Beispiel: genehmigt")
                save_absence = st.form_submit_button("Abwesenheit speichern", type="primary", width="stretch")

            if save_absence:
                if end_value < start_value:
                    st.warning("Das Enddatum darf nicht vor dem Startdatum liegen.")
                else:
                    add_vorgabe(
                        mitarbeitende_id=employee_options[selected_employee],
                        code=label_to_code[selected_label],
                        startdatum=start_value.isoformat(),
                        enddatum=end_value.isoformat(),
                        bemerkung=note.strip(),
                    )
                    invalidate_account_cache("Abwesenheiten")
                    st.success("Die Abwesenheit wurde gespeichert.")
                    st.rerun()

        with overview_tab:
            if not vorgaben:
                st.info("Noch keine Abwesenheiten gespeichert.")
            else:
                employee_names = {int(person["id"]): person["name"] for person in employees}
                overview_rows = []
                delete_options = {}
                for item in vorgaben:
                    person_name = employee_names.get(int(item["mitarbeitende_id"]), "Unbekannt")
                    label = VORGABE_LABELS.get(
                        str(item["code"]),
                        LEGACY_VORGABE_LABELS.get(str(item["code"]), str(item["code"])),
                    )
                    overview_rows.append(
                        {
                            "Mitarbeitende": person_name,
                            "Art": label,
                            "Von": german_date(item["startdatum"]),
                            "Bis": german_date(item["enddatum"]),
                            "Bemerkung": item.get("bemerkung", ""),
                        }
                    )
                    delete_options[
                        f'{person_name} · {label} · {german_date(item["startdatum"])} bis {german_date(item["enddatum"])}'
                    ] = int(item["id"])

                st.dataframe(
                    pd.DataFrame(overview_rows),
                    hide_index=True,
                    width="stretch",
                    column_config={
                        "Mitarbeitende": st.column_config.TextColumn("Mitarbeitende", width="medium"),
                        "Art": st.column_config.TextColumn("Art", width="small"),
                        "Von": st.column_config.TextColumn("Von", width="small"),
                        "Bis": st.column_config.TextColumn("Bis", width="small"),
                        "Bemerkung": st.column_config.TextColumn("Bemerkung", width="large"),
                    },
                )
                st.divider()
                selected_absence = st.selectbox("Eintrag auswählen", list(delete_options))
                if st.button("Ausgewählten Eintrag löschen", width="stretch"):
                    delete_vorgabe(delete_options[selected_absence])
                    invalidate_account_cache("Abwesenheiten")
                    st.rerun()


elif page == "Vorlagen":
    page_header(
        "Vorlagen",
        "Mit Vorlagen übernimmst du Besetzung und Planungsregeln. Eigene Einstellungen kannst du jederzeit speichern und wiederverwenden.",
    )

    section_title("Schnellvorlagen", "Eine Schnellvorlage ändert nur Besetzung und Planungsregeln. Mitarbeitende und Abwesenheiten bleiben erhalten.")
    template_columns = st.columns(3)
    for column, (template_name, template_data) in zip(template_columns, BUILTIN_TEMPLATES.items()):
        with column:
            values = template_data["settings"]
            st.markdown(
                f"""
                <div class="np-template-card">
                    <div class="np-template-title">{escape(template_name)}</div>
                    <div class="np-template-text">{escape(template_data['description'])}</div>
                    <hr style="margin:.75rem 0;">
                    <div class="np-template-text">
                        Früh {values['min_f']} · Spät {values['min_s']} · Nacht {values['min_n']}<br>
                        Max. {values['max_arbeitstage']} Arbeitstage · {values['max_wochenenden']} Wochenenden
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            if st.button(f'„{template_name}“ verwenden', key=f'builtin_{template_name}', width="stretch"):
                apply_template_settings(values)
                st.success(f"Die Vorlage „{template_name}“ wurde übernommen.")
                st.rerun()

    st.divider()
    own_tab, saved_tab = st.tabs(["💾 Eigene Vorlage speichern", "📚 Gespeicherte Vorlagen"])

    with own_tab:
        section_title("Aktuelle Einstellungen als Vorlage speichern")
        with st.form("save_template_form"):
            template_name = st.text_input("Name der Vorlage", placeholder="Zum Beispiel: Station C4 Standard")
            template_description = st.text_area(
                "Kurze Beschreibung",
                placeholder="Wann soll diese Vorlage verwendet werden?",
            )
            include_open_plan = st.checkbox(
                "Geöffneten Monatsplan ebenfalls speichern",
                value=False,
                disabled=not bool(st.session_state.dienstplan_vorlage),
                help="Damit kannst du später auch die Dienste nach Kalendertag in einen neuen Monat übernehmen.",
            )
            save_custom_template = st.form_submit_button("Vorlage speichern", type="primary", width="stretch")

        if save_custom_template:
            if not template_name.strip():
                st.warning("Bitte gib der Vorlage einen Namen.")
            else:
                payload = {
                    "settings": settings_snapshot(settings),
                    "plan": normalise_plan(st.session_state.dienstplan_vorlage) if include_open_plan else [],
                }
                save_vorlage(
                    name=template_name.strip(),
                    beschreibung=template_description.strip(),
                    kategorie="Monatsplan" if include_open_plan else "Besetzung & Regeln",
                    daten=payload,
                )
                invalidate_account_cache("Vorlagen")
                st.success("Die Vorlage wurde gespeichert.")
                st.rerun()

    with saved_tab:
        if not saved_templates:
            st.info("Noch keine eigenen Vorlagen gespeichert.")
        else:
            for template in saved_templates:
                with st.expander(f'{template["name"]} · {template["kategorie"]}'):
                    st.write(template.get("beschreibung") or "Keine Beschreibung hinterlegt.")
                    template_id = int(template["id"])
                    loaded_template = load_account_dataset(
                        f"Vorlagen/{template_id}",
                        lambda template_id=template_id: load_vorlage(template_id),
                    )
                    template_payload = loaded_template.get("daten", {}) if loaded_template else {}
                    settings_values = template_payload.get("settings", {})
                    plan_values = template_payload.get("plan", [])

                    if settings_values:
                        st.caption(
                            f'Besetzung: Früh {settings_values.get("min_f", "–")} · '
                            f'Spät {settings_values.get("min_s", "–")} · '
                            f'Nacht {settings_values.get("min_n", "–")}'
                        )
                    use_col, delete_col = st.columns(2)
                    with use_col:
                        if st.button("Einstellungen übernehmen", key=f'use_template_{template["id"]}', width="stretch"):
                            apply_template_settings(settings_values)
                            st.success("Die Einstellungen wurden übernommen.")
                            st.rerun()
                    with delete_col:
                        if st.button("Vorlage löschen", key=f'delete_template_{template["id"]}', width="stretch"):
                            delete_vorlage(int(template["id"]))
                            invalidate_account_cache("Vorlagen")
                            st.rerun()

                    if plan_values:
                        st.markdown("**Monatsplan in einen neuen Monat kopieren**")
                        c1, c2 = st.columns(2)
                        today = date.today()
                        with c1:
                            copy_month = st.selectbox(
                                "Zielmonat",
                                list(MONATSNAMEN),
                                index=today.month - 1,
                                format_func=lambda value: MONATSNAMEN[value],
                                key=f'copy_template_month_{template["id"]}',
                            )
                        with c2:
                            copy_year = st.number_input(
                                "Zieljahr",
                                min_value=2020,
                                max_value=2100,
                                value=today.year,
                                step=1,
                                key=f'copy_template_year_{template["id"]}',
                            )
                        if st.button("Monatsplan kopieren", key=f'copy_template_plan_{template["id"]}', type="primary", width="stretch"):
                            latest_vorgaben = get_vorgaben()
                            copied_plan = copy_plan_to_month(
                                plan_values,
                                int(copy_year),
                                int(copy_month),
                                employees,
                                latest_vorgaben,
                            )
                            invalidate_account_cache("Abwesenheiten")
                            st.session_state.dienstplan_vorlage = copied_plan
                            st.session_state.aktiver_dienstplan_id = None
                            st.session_state.dienstplan_name = f"Dienstplan {MONATSNAMEN[int(copy_month)]} {int(copy_year)}"
                            st.session_state.month_open_message = (
                                f"{count_absence_days(copied_plan)} Abwesenheitstage wurden übernommen."
                            )
                            reset_month_editor_draft()
                            st.success("Der Monatsplan wurde kopiert und geöffnet.")
                            st.rerun()


elif page == "Dienstplan":
    page_header(
        "Dienstplan erstellen",
        "Wähle einen Monat, übernimm Abwesenheiten und erstelle den Plan automatisch oder bearbeite ihn selbst.",
    )

    if not employees:
        st.warning("Bitte lege zuerst Mitarbeitende an.")
    else:
        new_tab, copy_tab, load_tab = st.tabs(["✨ Neuer Monat", "📋 Vorherigen Plan kopieren", "📂 Gespeicherten Plan laden"])

        today = date.today()
        with new_tab:
            c1, c2 = st.columns(2)
            with c1:
                selected_month = st.selectbox(
                    "Monat",
                    list(MONATSNAMEN),
                    index=today.month - 1,
                    format_func=lambda value: MONATSNAMEN[value],
                    key="new_plan_month",
                )
            with c2:
                selected_year = st.number_input(
                    "Jahr",
                    min_value=2020,
                    max_value=2100,
                    value=today.year,
                    step=1,
                    key="new_plan_year",
                )
            open_column, auto_column = st.columns(2)
            with open_column:
                if st.button("Leeren Monatsplan öffnen", width="stretch"):
                    latest_vorgaben = get_vorgaben()
                    prepared_plan = apply_vorgaben(
                        int(selected_year), int(selected_month), employees, latest_vorgaben
                    )
                    invalidate_account_cache("Abwesenheiten")
                    st.session_state.dienstplan_vorlage = prepared_plan
                    st.session_state.aktiver_dienstplan_id = None
                    st.session_state.dienstplan_name = f"Dienstplan {MONATSNAMEN[int(selected_month)]} {int(selected_year)}"
                    st.session_state.auto_plan_result = None
                    st.session_state.month_open_message = (
                        f"{count_absence_days(prepared_plan)} Abwesenheitstage wurden übernommen."
                    )
                    reset_month_editor_draft()
                    st.rerun()
            with auto_column:
                if st.button("Automatisch erstellen", type="primary", width="stretch"):
                    latest_vorgaben = get_vorgaben()
                    automatic_plan, automatic_result = create_automatic_month(
                        int(selected_year),
                        int(selected_month),
                        employees,
                        settings,
                        latest_vorgaben,
                    )
                    invalidate_account_cache("Abwesenheiten")
                    st.session_state.dienstplan_vorlage = automatic_plan
                    st.session_state.aktiver_dienstplan_id = None
                    st.session_state.dienstplan_name = f"Dienstplan {MONATSNAMEN[int(selected_month)]} {int(selected_year)}"
                    st.session_state.auto_plan_result = automatic_result
                    st.session_state.month_open_message = (
                        f"{count_absence_days(automatic_plan)} Abwesenheitstage wurden übernommen."
                    )
                    reset_month_editor_draft()
                    st.rerun()

        with copy_tab:
            if not saved_plans:
                st.info("Noch kein gespeicherter Dienstplan vorhanden.")
            else:
                saved_options = {
                    f'{plan["name"]} · {german_date(plan["startdatum"])}': int(plan["id"])
                    for plan in saved_plans
                }
                source_label = st.selectbox("Welcher Plan soll kopiert werden?", list(saved_options), key="copy_source_plan")
                c1, c2 = st.columns(2)
                with c1:
                    copy_month = st.selectbox(
                        "Zielmonat",
                        list(MONATSNAMEN),
                        index=today.month - 1,
                        format_func=lambda value: MONATSNAMEN[value],
                        key="copy_plan_month",
                    )
                with c2:
                    copy_year = st.number_input(
                        "Zieljahr",
                        min_value=2020,
                        max_value=2100,
                        value=today.year,
                        step=1,
                        key="copy_plan_year",
                    )
                st.caption("Die Dienste werden nach Kalendertag übernommen. Neue Abwesenheiten haben immer Vorrang.")
                if st.button("Plan kopieren und öffnen", type="primary", width="stretch"):
                    source_id = saved_options[source_label]
                    source = load_account_dataset(
                        f"Dienstpläne/{source_id}",
                        lambda source_id=source_id: load_dienstplan(source_id),
                    )
                    if source:
                        latest_vorgaben = get_vorgaben()
                        copied_plan = copy_plan_to_month(
                            source["plan"], int(copy_year), int(copy_month), employees, latest_vorgaben
                        )
                        invalidate_account_cache("Abwesenheiten")
                        st.session_state.dienstplan_vorlage = copied_plan
                        st.session_state.aktiver_dienstplan_id = None
                        st.session_state.dienstplan_name = f"Dienstplan {MONATSNAMEN[int(copy_month)]} {int(copy_year)}"
                        st.session_state.auto_plan_result = None
                        st.session_state.month_open_message = (
                            f"{count_absence_days(copied_plan)} Abwesenheitstage wurden übernommen."
                        )
                        reset_month_editor_draft()
                        st.rerun()

        with load_tab:
            if not saved_plans:
                st.info("Noch kein Plan gespeichert.")
            else:
                load_options = {
                    f'{plan["name"]} · {german_date(plan["startdatum"])}': int(plan["id"])
                    for plan in saved_plans
                }
                selected_saved = st.selectbox("Gespeicherten Plan auswählen", list(load_options), key="load_saved_plan")
                c1, c2 = st.columns(2)
                with c1:
                    if st.button("Plan öffnen", type="primary", width="stretch"):
                        selected_plan_id = load_options[selected_saved]
                        loaded = load_account_dataset(
                            f"Dienstpläne/{selected_plan_id}",
                            lambda selected_plan_id=selected_plan_id: load_dienstplan(selected_plan_id),
                        )
                        if loaded:
                            st.session_state.dienstplan_vorlage = normalise_plan(loaded["plan"])
                            st.session_state.aktiver_dienstplan_id = loaded["id"]
                            st.session_state.dienstplan_name = loaded["name"]
                            st.session_state.auto_plan_result = None
                            reset_month_editor_draft()
                            st.rerun()
                with c2:
                    if st.button("Gespeicherten Plan löschen", width="stretch"):
                        delete_dienstplan(load_options[selected_saved])
                        invalidate_account_cache("Dienstpläne")
                        st.rerun()

        if not st.session_state.dienstplan_vorlage:
            st.info("Öffne oben einen neuen Monat oder lade einen gespeicherten Plan.")
        else:
            st.divider()
            current_plan = st.session_state.dienstplan_vorlage
            year, month = month_from_plan(current_plan)
            current_errors, current_warnings = validate_plan(current_plan, employees, settings)
            current_stats = calculate_hours(current_plan, employees, settings)
            team_target = float(current_stats["Soll"].sum()) if not current_stats.empty else 0.0
            team_actual = float(current_stats["Ist"].sum()) if not current_stats.empty else 0.0

            month_open_message = st.session_state.pop("month_open_message", None)
            if month_open_message:
                st.success(month_open_message)

            c1, c2, c3, c4 = st.columns(4)
            with c1:
                kpi_card("Geöffneter Monat", f"{MONATSNAMEN[month]} {year}", st.session_state.dienstplan_name)
            with c2:
                kpi_card("Geplante Stunden", f"{team_actual:.1f}", f"von {team_target:.1f} Sollstunden")
            with c3:
                kpi_card("Fehler", str(len(current_errors)), "müssen vor Freigabe geprüft werden")
            with c4:
                kpi_card("Hinweise", str(len(current_warnings)), "zum Beispiel Unterbesetzung")

            show_plan_legend()

            auto_result = st.session_state.get("auto_plan_result")
            if auto_result:
                shortages = auto_result.get("shortages", [])
                st.success(
                    f'Automatisch eingetragen: {int(auto_result.get("assigned", 0))} Dienste. '
                    "Gespeicherte Abwesenheiten wurden nicht überschrieben."
                )
                if shortages:
                    with st.expander(f"Offene Besetzungen: {len(shortages)}"):
                        for message in shortages[:50]:
                            st.warning(message)

            edit_plan_tab, readable_tab, check_tab, save_tab = st.tabs(
                ["✏️ Plan bearbeiten", "👀 Lesbare Ansicht", "✅ Prüfung und Stunden", "💾 Speichern und Export"]
            )

            preview_plan = current_plan
            preview_stats = current_stats
            over_target = [
                row for _, row in preview_stats.iterrows()
                if float(row["Ist"]) > float(row["Soll"]) + 0.05
            ]

            with edit_plan_tab:
                st.markdown("#### Dienste mit dem Pinsel eintragen")
                st.caption(
                    "Wähle oben F, S, N, Frei oder den Radierer. Klicke anschließend auf "
                    "ein Feld oder ziehe mit gedrückter Maustaste über mehrere Felder. "
                    "Die Seite lädt erst bei „Dienste übernehmen“ neu."
                )

                saved_message = st.session_state.pop("manual_month_saved_message", None)
                if saved_message:
                    st.success(saved_message)

                paint_days, paint_rows = plan_to_paint_grid(current_plan, employees, settings)
                if paint_grid_component is None:
                    st.error(
                        "Der Ordner paint_grid/dist fehlt. Lade den vollständigen "
                        "Ordner paint_grid zusammen mit app.py auf GitHub hoch."
                    )
                    paint_result = None
                else:
                    paint_result = paint_grid_component(
                        days=paint_days,
                        rows=paint_rows,
                        shift_hours={
                            "F": float(settings.get("f_stunden", 0.0)),
                            "S": float(settings.get("s_stunden", 0.0)),
                            "N": float(settings.get("n_stunden", 0.0)),
                        },
                        plan_version=(
                            f'{st.session_state.dienstplan_editor_version}:'
                            f'{year:04d}-{month:02d}'
                        ),
                        key=f"paint_grid_{st.session_state.dienstplan_editor_version}",
                        default=None,
                    )

                st.caption(
                    "U, W, FB und K kommen aus den Abwesenheiten und bleiben gesperrt. "
                    "Soll-, Ist- und Reststunden stehen direkt neben den Namen und werden "
                    "beim Eintragen sofort mitgerechnet."
                )

                if isinstance(paint_result, dict):
                    result_token = str(paint_result.get("token", ""))
                    if (
                        result_token
                        and result_token != st.session_state.get("paint_grid_last_token")
                    ):
                        st.session_state.paint_grid_last_token = result_token
                        edited_matrix = paint_result_to_matrix(
                            paint_result,
                            current_plan,
                            employees,
                            settings,
                        )
                        input_errors = manual_month_input_errors(
                            edited_matrix,
                            current_plan,
                            employees,
                        )
                        candidate_plan = matrix_to_plan(edited_matrix, current_plan)
                        candidate_stats = calculate_hours(candidate_plan, employees, settings)
                        candidate_over_target = [
                            row for _, row in candidate_stats.iterrows()
                            if float(row["Ist"]) > float(row["Soll"]) + 0.05
                        ]
                        if input_errors:
                            for message in input_errors:
                                st.error(message)
                        elif candidate_over_target:
                            names = ", ".join(
                                str(row["Mitarbeitende"])
                                for row in candidate_over_target
                            )
                            st.error(
                                f"Noch nicht übernommen: {names} wäre über den Sollstunden."
                            )
                        else:
                            st.session_state.dienstplan_vorlage = candidate_plan
                            st.session_state.auto_plan_result = None
                            st.session_state.manual_month_saved_message = (
                                "Die Dienste wurden übernommen. Stunden und Prüfungen sind aktualisiert."
                            )
                            reset_month_editor_draft()
                            st.rerun()

            with readable_tab:
                st.caption("Diese Ansicht ist nur zum Lesen. Die Farben erleichtern die Kontrolle der Dienste.")
                st.dataframe(
                    readable_plan_matrix(plan_to_matrix(preview_plan, employees, settings)),
                    hide_index=True,
                    width="stretch",
                    height=min(760, 105 + len(employees) * 38),
                )

            errors, warnings = validate_plan(preview_plan, employees, settings)
            with check_tab:
                stats_display = preview_stats.rename(
                    columns={
                        "Soll": "Sollstunden",
                        "Ist": "Geplant",
                        "+/-": "Abweichung",
                        "F": "Früh",
                        "S": "Spät",
                        "N": "Nacht",
                        "U": "Urlaub",
                        "W": "Wunschfrei",
                        "FB": "Fortbildung",
                        "K": "Krank",
                    }
                )
                st.dataframe(stats_display, hide_index=True, width="stretch")

                if errors:
                    with st.expander(f"Fehler: {len(errors)}", expanded=True):
                        for message in errors:
                            st.error(message)
                else:
                    st.success("Keine Fehler gefunden.")

                if warnings:
                    with st.expander(f"Hinweise: {len(warnings)}"):
                        for message in warnings[:60]:
                            st.warning(message)
                        if len(warnings) > 60:
                            st.caption(f"Weitere {len(warnings) - 60} Hinweise werden nicht angezeigt.")

            with save_tab:
                plan_name = st.text_input("Name des Dienstplans", value=st.session_state.dienstplan_name)
                if over_target:
                    st.error("Mindestens eine Person liegt über den Sollstunden. Entferne zuerst einen Dienst.")
                if st.button("Dienstplan speichern", type="primary", width="stretch", disabled=bool(over_target)):
                    plan_id = save_dienstplan(
                        name=plan_name.strip() or "Dienstplan",
                        startdatum=preview_plan[0]["Datum"],
                        zeitraum=len(preview_plan),
                        plan=preview_plan,
                        plan_id=st.session_state.aktiver_dienstplan_id,
                    )
                    invalidate_account_cache("Dienstpläne")
                    st.session_state.aktiver_dienstplan_id = plan_id
                    st.session_state.dienstplan_name = plan_name.strip() or "Dienstplan"
                    st.session_state.dienstplan_vorlage = preview_plan
                    st.success("Der Dienstplan wurde gespeichert.")

                export_name = safe_filename(plan_name)
                export_1, export_2 = st.columns(2)
                with export_1:
                    st.download_button(
                        "Als Excel herunterladen",
                        data=create_excel(preview_plan, plan_name, employees, settings),
                        file_name=f"{export_name}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        width="stretch",
                        disabled=bool(over_target),
                    )
                with export_2:
                    st.download_button(
                        "Als PDF herunterladen",
                        data=create_pdf(preview_plan, plan_name, employees, settings),
                        file_name=f"{export_name}.pdf",
                        mime="application/pdf",
                        width="stretch",
                        disabled=bool(over_target),
                    )


elif page == "Auswertung":
    page_header(
        "Auswertung",
        "Vergleiche Soll- und Iststunden, Schichtverteilung und Planungsfehler für den geöffneten Dienstplan.",
    )

    if not st.session_state.dienstplan_vorlage:
        st.info("Öffne zuerst einen Dienstplan auf der Seite „Dienstplan“.")
    elif not employees:
        st.info("Es sind keine Mitarbeitenden vorhanden.")
    else:
        plan = st.session_state.dienstplan_vorlage
        stats = calculate_hours(plan, employees, settings)
        errors, warnings = validate_plan(plan, employees, settings)
        target_total = float(stats["Soll"].sum()) if not stats.empty else 0.0
        actual_total = float(stats["Ist"].sum()) if not stats.empty else 0.0
        remaining_total = target_total - actual_total

        c1, c2, c3, c4 = st.columns(4)
        with c1:
            kpi_card("Sollstunden", f"{target_total:.1f}", "gesamtes Team")
        with c2:
            kpi_card("Geplante Stunden", f"{actual_total:.1f}", "inklusive bezahlter Abwesenheit")
        with c3:
            kpi_card("Noch offen", f"{remaining_total:.1f}", "Stunden bis zum Teamsoll")
        with c4:
            kpi_card("Planprüfung", f"{len(errors)} / {len(warnings)}", "Fehler / Hinweise")

        overview_tab, chart_tab, quality_tab = st.tabs(["Stundenübersicht", "Schichtverteilung", "Planqualität"])
        with overview_tab:
            display = stats.rename(
                columns={
                    "Soll": "Sollstunden",
                    "Ist": "Geplant",
                    "+/-": "Abweichung",
                    "F": "Früh",
                    "S": "Spät",
                    "N": "Nacht",
                    "U": "Urlaub",
                    "W": "Wunschfrei",
                    "FB": "Fortbildung",
                    "K": "Krank",
                }
            )
            st.dataframe(display, hide_index=True, width="stretch")

        with chart_tab:
            chart_data = stats.set_index("Mitarbeitende")[["F", "S", "N"]].rename(
                columns={"F": "Frühdienst", "S": "Spätdienst", "N": "Nachtdienst"}
            )
            st.bar_chart(chart_data)
            st.caption("Die Grafik zeigt die Anzahl der geplanten Früh-, Spät- und Nachtdienste je Person.")

        with quality_tab:
            if errors:
                section_title(f"Fehler ({len(errors)})")
                for message in errors:
                    st.error(message)
            else:
                st.success("Keine Fehler gefunden.")
            if warnings:
                section_title(f"Hinweise ({len(warnings)})")
                for message in warnings[:60]:
                    st.warning(message)


elif page == "Planungsregeln":
    page_header(
        "Planungsregeln",
        "Bestimme selbst, welche Regeln die automatische Planung einhalten soll. Jede Änderung gilt nur für dein Konto.",
    )

    custom_rules = eigene_planungsregeln(settings)
    active_basic_count = sum(
        1
        for key in [
            "regel_ruhezeit_aktiv",
            "regel_max_arbeitstage_aktiv",
            "regel_max_nachtdienste_aktiv",
            "regel_max_wochenenden_aktiv",
            "keine_folgewochenenden",
            "regel_sollstunden_aktiv",
            "regel_fachkraft_aktiv",
            "regel_nacht_zuerst_aktiv",
            "regel_nachtblock_aktiv",
            "regel_wochenende_zuerst_aktiv",
            "regel_wochenende_gleich_aktiv",
        ]
        if bool(settings.get(key, True))
    )
    active_custom_count = sum(1 for rule in custom_rules if bool(rule.get("aktiv", True)))

    c1, c2, c3 = st.columns(3)
    with c1:
        kpi_card("Aktive Grundregeln", str(active_basic_count), "für Schutz und Planungsreihenfolge")
    with c2:
        kpi_card("Eigene Regeln", str(len(custom_rules)), f"davon {active_custom_count} aktiv")
    with c3:
        kpi_card("Speicherung", "Automatisch", "persönlich im Benutzerkonto")

    basic_tab, custom_tab, overview_tab = st.tabs(
        ["Grundregeln", "Eigene Wenn–Dann-Regeln", "Regelübersicht"]
    )

    with basic_tab:
        section_title(
            "Schutz- und Belastungsregeln",
            "Ein Häkchen aktiviert die Regel. Der Wert daneben bestimmt die Grenze.",
        )
        with st.form("basic_planning_rules_form"):
            left, right = st.columns(2, gap="large")
            with left:
                rest_active = st.checkbox(
                    "Mindestruhezeit einhalten",
                    value=regel_aktiv(settings, "regel_ruhezeit_aktiv"),
                )
                rest_hours = st.number_input(
                    "Mindestruhezeit in Stunden",
                    min_value=0.0,
                    max_value=24.0,
                    value=float(settings.get("ruhezeit", 11.0)),
                    step=0.5,
                )
                max_work_active = st.checkbox(
                    "Arbeitstage hintereinander begrenzen",
                    value=regel_aktiv(settings, "regel_max_arbeitstage_aktiv"),
                )
                max_work_days = st.number_input(
                    "Maximale Arbeitstage hintereinander",
                    min_value=1,
                    max_value=14,
                    value=int(settings.get("max_arbeitstage", 6)),
                    step=1,
                )
                target_active = st.checkbox(
                    "Persönliche Sollstunden nicht überschreiten",
                    value=regel_aktiv(settings, "regel_sollstunden_aktiv"),
                )
            with right:
                max_nights_active = st.checkbox(
                    "Nachtdienste hintereinander begrenzen",
                    value=regel_aktiv(settings, "regel_max_nachtdienste_aktiv"),
                )
                max_night_days = st.number_input(
                    "Maximale Nachtdienste hintereinander",
                    min_value=1,
                    max_value=10,
                    value=int(settings.get("max_nachtdienste", 4)),
                    step=1,
                )
                max_weekends_active = st.checkbox(
                    "Arbeitswochenenden pro Monat begrenzen",
                    value=regel_aktiv(settings, "regel_max_wochenenden_aktiv"),
                )
                max_weekends = st.number_input(
                    "Maximale Arbeitswochenenden pro Monat",
                    min_value=0,
                    max_value=5,
                    value=int(settings.get("max_wochenenden", 2)),
                    step=1,
                )
                no_consecutive_weekends = st.checkbox(
                    "Keine direkt aufeinanderfolgenden Arbeitswochenenden",
                    value=bool(settings.get("keine_folgewochenenden", True)),
                )

            st.divider()
            section_title(
                "Qualität und Reihenfolge",
                "Diese Regeln steuern Besetzung und Ablauf der automatischen Planung.",
            )
            q1, q2 = st.columns(2, gap="large")
            with q1:
                fachkraft_active = st.checkbox(
                    "Mindestzahl an Fachkräften einhalten",
                    value=regel_aktiv(settings, "regel_fachkraft_aktiv"),
                )
                night_first_active = st.checkbox(
                    "Zuerst alle Nachtdienste planen",
                    value=regel_aktiv(settings, "regel_nacht_zuerst_aktiv"),
                )
                night_block_active = st.checkbox(
                    "Begonnene Nachtblöcke möglichst fortsetzen",
                    value=regel_aktiv(settings, "regel_nachtblock_aktiv"),
                )
            with q2:
                weekend_first_active = st.checkbox(
                    "Wochenenden vor den Werktagen planen",
                    value=regel_aktiv(settings, "regel_wochenende_zuerst_aktiv"),
                )
                same_weekend_active = st.checkbox(
                    "Samstag und Sonntag möglichst gleicher Dienst",
                    value=regel_aktiv(settings, "regel_wochenende_gleich_aktiv"),
                )

            save_basic_rules = st.form_submit_button(
                "Grundregeln speichern", type="primary", width="stretch"
            )

        if save_basic_rules:
            new_settings = dict(settings)
            new_settings.update(
                {
                    "regel_ruhezeit_aktiv": bool(rest_active),
                    "ruhezeit": float(rest_hours),
                    "regel_max_arbeitstage_aktiv": bool(max_work_active),
                    "max_arbeitstage": int(max_work_days),
                    "regel_sollstunden_aktiv": bool(target_active),
                    "regel_max_nachtdienste_aktiv": bool(max_nights_active),
                    "max_nachtdienste": int(max_night_days),
                    "regel_max_wochenenden_aktiv": bool(max_weekends_active),
                    "max_wochenenden": int(max_weekends),
                    "keine_folgewochenenden": bool(no_consecutive_weekends),
                    "regel_fachkraft_aktiv": bool(fachkraft_active),
                    "regel_nacht_zuerst_aktiv": bool(night_first_active),
                    "regel_nachtblock_aktiv": bool(night_block_active),
                    "regel_wochenende_zuerst_aktiv": bool(weekend_first_active),
                    "regel_wochenende_gleich_aktiv": bool(same_weekend_active),
                }
            )
            save_einstellungen(new_settings)
            st.session_state.settings = new_settings
            st.success("Die Grundregeln wurden gespeichert.")
            st.rerun()

        st.info(
            "Die Mindestbesetzung pro Schicht wird weiterhin unter „Einstellungen → Mindestbesetzung“ festgelegt. "
            "Tarifvertrag, Betriebsvereinbarung und interne Vorgaben müssen zusätzlich geprüft werden."
        )

    with custom_tab:
        section_title(
            "Eigene Regel erstellen oder bearbeiten",
            "Wähle eine Regelart aus und fülle die angezeigten Auswahlfelder aus. Programmieren ist nicht nötig.",
        )
        if not employees:
            st.warning("Lege zuerst mindestens eine mitarbeitende Person an.")
        else:
            rule_type_labels = {
                "nicht_gemeinsam": "Zwei Personen nicht gemeinsam einplanen",
                "frei_nach_dienst": "Freie Tage nach einer Dienstart",
                "max_dienstfolge": "Dienstart hintereinander begrenzen",
                "nicht_an_wochentagen": "Person kann an bestimmten Wochentagen nicht arbeiten",
            }
            edit_options = {"Neue Regel erstellen": "NEW"}
            for rule in custom_rules:
                edit_options[f'Bearbeiten: {rule.get("name", "Eigene Regel")}'] = str(rule.get("id"))

            selected_label = st.selectbox(
                "Neue oder vorhandene Regel",
                list(edit_options),
                key="custom_rule_editor_selection",
            )
            selected_rule_id = edit_options[selected_label]
            selected_rule = next(
                (rule for rule in custom_rules if str(rule.get("id")) == selected_rule_id),
                None,
            )
            selected_type = str((selected_rule or {}).get("typ", "nicht_gemeinsam"))
            type_names = list(rule_type_labels)
            rule_type = st.selectbox(
                "Regelart",
                type_names,
                index=type_names.index(selected_type) if selected_type in type_names else 0,
                format_func=lambda value: rule_type_labels[value],
                key=f"custom_rule_type_{selected_rule_id}",
            )

            employee_ids = [int(person["id"]) for person in employees]
            employee_by_id_ui = {int(person["id"]): person for person in employees}
            shift_labels = {"F": "Frühdienst", "S": "Spätdienst", "N": "Nachtdienst"}
            editor_token = f"{selected_rule_id}_{rule_type}"

            with st.form(f"custom_rule_form_{editor_token}"):
                rule_name = st.text_input(
                    "Name der Regel",
                    value=str((selected_rule or {}).get("name", "")),
                    placeholder="Zum Beispiel: Ilona und Marion nicht zusammen",
                    max_chars=100,
                )
                rule_enabled = st.checkbox(
                    "Regel sofort aktivieren",
                    value=bool((selected_rule or {}).get("aktiv", True)),
                )

                if rule_type == "nicht_gemeinsam":
                    first_default = regel_zahl(selected_rule or {}, "mitarbeitende_a_id", employee_ids[0])
                    second_default = regel_zahl(
                        selected_rule or {},
                        "mitarbeitende_b_id",
                        employee_ids[1] if len(employee_ids) > 1 else employee_ids[0],
                    )
                    first_id = st.selectbox(
                        "Erste Person",
                        employee_ids,
                        index=employee_ids.index(first_default) if first_default in employee_ids else 0,
                        format_func=lambda value: employee_by_id_ui[value]["name"],
                    )
                    second_id = st.selectbox(
                        "Zweite Person",
                        employee_ids,
                        index=employee_ids.index(second_default) if second_default in employee_ids else 0,
                        format_func=lambda value: employee_by_id_ui[value]["name"],
                    )
                    together_shift_options = ["ALL", "F", "S", "N"]
                    together_shift_labels = {"ALL": "Alle Dienstarten", **shift_labels}
                    old_shift = str((selected_rule or {}).get("dienst", "ALL"))
                    selected_shift = st.selectbox(
                        "Gilt für",
                        together_shift_options,
                        index=together_shift_options.index(old_shift) if old_shift in together_shift_options else 0,
                        format_func=lambda value: together_shift_labels[value],
                    )

                elif rule_type == "frei_nach_dienst":
                    scope_ids = [0] + employee_ids
                    scope_default = regel_zahl(selected_rule or {}, "mitarbeitende_id", 0)
                    selected_employee_id = st.selectbox(
                        "Gilt für",
                        scope_ids,
                        index=scope_ids.index(scope_default) if scope_default in scope_ids else 0,
                        format_func=lambda value: "Alle Mitarbeitenden" if value == 0 else employee_by_id_ui[value]["name"],
                    )
                    old_shift = str((selected_rule or {}).get("dienst", "N"))
                    selected_shift = st.selectbox(
                        "Nach welcher Dienstart?",
                        list(shift_labels),
                        index=list(shift_labels).index(old_shift) if old_shift in shift_labels else 2,
                        format_func=lambda value: shift_labels[value],
                    )
                    free_days_value = st.number_input(
                        "Mindestens freie Kalendertage danach",
                        min_value=1,
                        max_value=7,
                        value=max(1, regel_zahl(selected_rule or {}, "freie_tage", 2)),
                        step=1,
                    )

                elif rule_type == "max_dienstfolge":
                    scope_ids = [0] + employee_ids
                    scope_default = regel_zahl(selected_rule or {}, "mitarbeitende_id", 0)
                    selected_employee_id = st.selectbox(
                        "Gilt für",
                        scope_ids,
                        index=scope_ids.index(scope_default) if scope_default in scope_ids else 0,
                        format_func=lambda value: "Alle Mitarbeitenden" if value == 0 else employee_by_id_ui[value]["name"],
                    )
                    old_shift = str((selected_rule or {}).get("dienst", "N"))
                    selected_shift = st.selectbox(
                        "Welche Dienstart?",
                        list(shift_labels),
                        index=list(shift_labels).index(old_shift) if old_shift in shift_labels else 2,
                        format_func=lambda value: shift_labels[value],
                    )
                    maximum_value = st.number_input(
                        "Höchstens hintereinander",
                        min_value=1,
                        max_value=10,
                        value=max(1, regel_zahl(selected_rule or {}, "maximum", 3)),
                        step=1,
                    )

                else:
                    scope_ids = [0] + employee_ids
                    scope_default = regel_zahl(selected_rule or {}, "mitarbeitende_id", 0)
                    selected_employee_id = st.selectbox(
                        "Welche Person?",
                        scope_ids,
                        index=scope_ids.index(scope_default) if scope_default in scope_ids else 0,
                        format_func=lambda value: "Alle Mitarbeitenden" if value == 0 else employee_by_id_ui[value]["name"],
                    )
                    selected_weekdays = st.multiselect(
                        "An welchen Wochentagen kann die Person nicht arbeiten?",
                        list(range(7)),
                        default=sorted(regel_wochentage(selected_rule or {})),
                        format_func=lambda value: [
                            "Montag", "Dienstag", "Mittwoch", "Donnerstag",
                            "Freitag", "Samstag", "Sonntag",
                        ][value],
                        placeholder="Wochentage auswählen",
                    )

                save_custom_rule = st.form_submit_button(
                    "Regel speichern", type="primary", width="stretch",
                    disabled=rule_type == "nicht_gemeinsam" and len(employee_ids) < 2,
                )

            if save_custom_rule:
                validation_message = ""
                if not rule_name.strip():
                    validation_message = "Bitte gib der Regel einen Namen."
                elif rule_type == "nicht_gemeinsam" and first_id == second_id:
                    validation_message = "Bitte wähle zwei unterschiedliche Personen."
                elif rule_type == "nicht_an_wochentagen" and not selected_weekdays:
                    validation_message = "Bitte wähle mindestens einen Wochentag aus."

                if validation_message:
                    st.warning(validation_message)
                else:
                    saved_rule = {
                        "id": str((selected_rule or {}).get("id") or uuid4()),
                        "name": rule_name.strip(),
                        "typ": rule_type,
                        "aktiv": bool(rule_enabled),
                    }
                    if rule_type == "nicht_gemeinsam":
                        saved_rule.update(
                            {
                                "mitarbeitende_a_id": int(first_id),
                                "mitarbeitende_b_id": int(second_id),
                                "dienst": selected_shift,
                            }
                        )
                    elif rule_type == "frei_nach_dienst":
                        saved_rule.update(
                            {
                                "mitarbeitende_id": int(selected_employee_id),
                                "dienst": selected_shift,
                                "freie_tage": int(free_days_value),
                            }
                        )
                    elif rule_type == "max_dienstfolge":
                        saved_rule.update(
                            {
                                "mitarbeitende_id": int(selected_employee_id),
                                "dienst": selected_shift,
                                "maximum": int(maximum_value),
                            }
                        )
                    else:
                        saved_rule.update(
                            {
                                "mitarbeitende_id": int(selected_employee_id),
                                "wochentage": sorted(int(value) for value in selected_weekdays),
                            }
                        )

                    updated_rules = [
                        rule
                        for rule in custom_rules
                        if str(rule.get("id")) != str(saved_rule["id"])
                    ]
                    updated_rules.append(saved_rule)
                    new_settings = {**settings, "eigene_planungsregeln": updated_rules}
                    save_einstellungen(new_settings)
                    st.session_state.settings = new_settings
                    st.success("Die eigene Regel wurde gespeichert.")
                    st.rerun()

            st.divider()
            section_title(
                "Gespeicherte eigene Regeln",
                "Deaktivierte Regeln bleiben gespeichert, werden aber nicht angewendet.",
            )
            if not custom_rules:
                st.info("Noch keine eigene Regel gespeichert.")
            else:
                for rule in custom_rules:
                    rule_id = str(rule.get("id"))
                    enabled = bool(rule.get("aktiv", True))
                    status_icon = "🟢" if enabled else "⚪"
                    with st.expander(f'{status_icon} {rule.get("name", "Eigene Regel")}'):
                        st.write(regel_beschreibung(rule, employees))
                        st.caption("Aktiv" if enabled else "Inaktiv")
                        action_col, delete_col = st.columns(2)
                        with action_col:
                            action_label = "Deaktivieren" if enabled else "Aktivieren"
                            if st.button(action_label, key=f"toggle_rule_{rule_id}", width="stretch"):
                                updated_rules = []
                                for item in custom_rules:
                                    updated = dict(item)
                                    if str(updated.get("id")) == rule_id:
                                        updated["aktiv"] = not enabled
                                    updated_rules.append(updated)
                                new_settings = {**settings, "eigene_planungsregeln": updated_rules}
                                save_einstellungen(new_settings)
                                st.session_state.settings = new_settings
                                st.rerun()
                        with delete_col:
                            if st.button("Löschen", key=f"delete_rule_{rule_id}", width="stretch"):
                                updated_rules = [
                                    item for item in custom_rules
                                    if str(item.get("id")) != rule_id
                                ]
                                new_settings = {**settings, "eigene_planungsregeln": updated_rules}
                                save_einstellungen(new_settings)
                                st.session_state.settings = new_settings
                                st.rerun()

    with overview_tab:
        section_title(
            "Alle Regeln auf einen Blick",
            "Nur aktive Regeln werden bei der automatischen Planung und Planprüfung angewendet.",
        )
        basic_rows = [
            ("Mindestruhezeit", "regel_ruhezeit_aktiv", f'{float(settings.get("ruhezeit", 11.0)):g} Stunden'),
            ("Maximale Arbeitstage", "regel_max_arbeitstage_aktiv", str(int(settings.get("max_arbeitstage", 6)))),
            ("Maximale Nachtdienste", "regel_max_nachtdienste_aktiv", str(int(settings.get("max_nachtdienste", 4)))),
            ("Maximale Arbeitswochenenden", "regel_max_wochenenden_aktiv", str(int(settings.get("max_wochenenden", 2)))),
            ("Keine Folgewochenenden", "keine_folgewochenenden", "–"),
            ("Sollstunden nicht überschreiten", "regel_sollstunden_aktiv", "–"),
            ("Fachkraftbesetzung", "regel_fachkraft_aktiv", "–"),
            ("Nachtdienste zuerst", "regel_nacht_zuerst_aktiv", "–"),
            ("Nachtblöcke fortsetzen", "regel_nachtblock_aktiv", "–"),
            ("Wochenenden zuerst", "regel_wochenende_zuerst_aktiv", "–"),
            ("Gleicher Wochenenddienst", "regel_wochenende_gleich_aktiv", "–"),
        ]
        overview_rows = [
            {
                "Regel": name,
                "Art": "Grundregel",
                "Status": "Aktiv" if bool(settings.get(key, True)) else "Inaktiv",
                "Wert": value,
            }
            for name, key, value in basic_rows
        ]
        overview_rows.extend(
            {
                "Regel": str(rule.get("name", "Eigene Regel")),
                "Art": "Eigene Regel",
                "Status": "Aktiv" if bool(rule.get("aktiv", True)) else "Inaktiv",
                "Wert": regel_beschreibung(rule, employees),
            }
            for rule in custom_rules
        )
        st.dataframe(
            pd.DataFrame(overview_rows),
            hide_index=True,
            width="stretch",
            column_config={
                "Regel": st.column_config.TextColumn("Regel", width="medium"),
                "Art": st.column_config.TextColumn("Art", width="small"),
                "Status": st.column_config.TextColumn("Status", width="small"),
                "Wert": st.column_config.TextColumn("Wert / Beschreibung", width="large"),
            },
        )


elif page == "Feedback & Hilfe":
    page_header(
        "Kontakt",
        "Hast du eine Frage, eine Idee oder ein Problem? Schreib uns einfach eine Nachricht.",
    )

    try:
        get_feedback(limit=1)
        feedback_available = True
    except Exception:
        feedback_available = False

    if not feedback_available:
        st.warning(
            "Die Feedback-Speicherung ist noch nicht aktiviert. Führe zuerst einmal "
            "die Datei „feedback_setup.sql“ im Supabase SQL Editor aus."
        )

    _, contact_column, _ = st.columns([1, 1.6, 1])
    with contact_column:
        with st.container(border=True):
            section_title(
                "Nachricht senden",
                "Wir lesen jede Nachricht und nutzen sie, um NursePlan Pro zu verbessern.",
            )
            with st.form("feedback_form", clear_on_submit=True):
                feedback_subject = st.text_input(
                    "Betreff (optional)",
                    placeholder="Worum geht es?",
                    max_chars=120,
                )
                feedback_message = st.text_area(
                    "Deine Nachricht",
                    placeholder="Schreib uns kurz, was du mitteilen möchtest …",
                    height=180,
                    max_chars=3000,
                )
                st.caption(
                    "Bitte keine Patientendaten, Passwörter oder Zugangsschlüssel eintragen."
                )
                submit_feedback = st.form_submit_button(
                    "Nachricht senden",
                    type="primary",
                    width="stretch",
                    disabled=not feedback_available,
                )

            if submit_feedback:
                if len(feedback_message.strip()) < 10:
                    st.warning("Bitte schreibe eine kurze Nachricht mit mindestens 10 Zeichen.")
                else:
                    try:
                        add_feedback(
                            kategorie="Frage",
                            bewertung=5,
                            betreff=feedback_subject.strip() or "Allgemeine Rückmeldung",
                            nachricht=feedback_message,
                            kontakt_erlaubt=True,
                            kontakt_email=current_user_email(),
                        )
                        st.success("Danke! Deine Nachricht wurde gesendet.")
                    except Exception:
                        st.error(
                            "Die Nachricht konnte gerade nicht gesendet werden. "
                            "Bitte versuche es später noch einmal."
                        )

        st.caption(
            "🔒 Deine Nachricht ist privat. Wenn eine Antwort nötig ist, nutzen wir die "
            "E-Mail-Adresse deines Kontos."
        )


elif page == "Rechtliches":
    render_legal_page()


elif page == "Einstellungen":
    page_header(
        "Einstellungen",
        "Lege Arbeitszeiten, Mindestbesetzung und Planungsgrenzen fest. Diese Werte gelten für neue automatische Dienstpläne.",
    )

    general_tab, shifts_tab, staffing_tab = st.tabs(
        ["Allgemein", "Arbeitszeiten", "Mindestbesetzung"]
    )

    with general_tab:
        section_title("Monatssoll", "Dieser Wert dient zur Berechnung des Beschäftigungsumfangs.")
        fulltime_target = st.number_input(
            "Vollzeit-Sollstunden im Monat",
            min_value=1.0,
            max_value=300.0,
            value=float(settings["vollzeit_monatssoll"]),
            step=0.5,
        )

    shift_values = {}
    with shifts_tab:
        section_title("Arbeitszeiten", "Beginn, Ende und bezahlte Stunden werden für Ruhezeit und Stundenkonto benötigt.")
        for code, prefix, title in [
            ("F", "f", "Frühdienst"),
            ("S", "s", "Spätdienst"),
            ("N", "n", "Nachtdienst"),
        ]:
            st.markdown(f"### {code} · {title}")
            c1, c2, c3 = st.columns(3)
            with c1:
                start_value = st.time_input(
                    "Beginn",
                    value=parse_clock(settings[f"{prefix}_start"]),
                    key=f"settings_{prefix}_start_new",
                )
            with c2:
                end_value = st.time_input(
                    "Ende",
                    value=parse_clock(settings[f"{prefix}_ende"]),
                    key=f"settings_{prefix}_end_new",
                )
            with c3:
                hours_value = st.number_input(
                    "Bezahlte Stunden",
                    min_value=0.0,
                    max_value=24.0,
                    value=float(settings[f"{prefix}_stunden"]),
                    step=0.1,
                    key=f"settings_{prefix}_hours_new",
                )
            shift_values[prefix] = {
                "start": format_clock(start_value),
                "end": format_clock(end_value),
                "hours": float(hours_value),
            }
            st.divider()

    staffing_values = {}
    with staffing_tab:
        section_title("Mindestbesetzung je Dienst", "Die Automatik versucht diese Besetzung an jedem Kalendertag zu erreichen.")
        for code, prefix, title in [
            ("F", "f", "Frühdienst"),
            ("S", "s", "Spätdienst"),
            ("N", "n", "Nachtdienst"),
        ]:
            c1, c2 = st.columns(2)
            with c1:
                minimum_value = st.number_input(
                    f"{title}: Personen insgesamt",
                    min_value=0,
                    max_value=30,
                    value=int(settings[f"min_{prefix}"]),
                    step=1,
                    key=f"settings_{prefix}_minimum_new",
                )
            with c2:
                fach_value = st.number_input(
                    f"{title}: davon Fachkräfte",
                    min_value=0,
                    max_value=30,
                    value=int(settings.get(f"min_fach_{prefix}", 1)),
                    step=1,
                    key=f"settings_{prefix}_fach_new",
                )
            staffing_values[prefix] = {"minimum": int(minimum_value), "fach": int(fach_value)}

    if st.button("Alle Einstellungen speichern", type="primary", width="stretch"):
        new_settings = dict(settings)
        new_settings["vollzeit_monatssoll"] = float(fulltime_target)
        new_settings["max_plusstunden"] = 0.0
        for prefix, values in shift_values.items():
            new_settings[f"{prefix}_start"] = values["start"]
            new_settings[f"{prefix}_ende"] = values["end"]
            new_settings[f"{prefix}_stunden"] = values["hours"]
        for prefix, values in staffing_values.items():
            new_settings[f"min_{prefix}"] = values["minimum"]
            new_settings[f"min_fach_{prefix}"] = values["fach"]

        save_einstellungen(new_settings)
        st.session_state.settings = new_settings
        st.success("Die Einstellungen wurden gespeichert.")
        st.rerun()

    st.divider()
    with st.expander("⚡ Leistung und Ladezeiten"):
        st.caption(
            "Die Messwerte gelten nur für deine aktuelle Sitzung. „Cache“ bedeutet, "
            "dass keine neue Datenbankabfrage nötig war."
        )
        metrics = st.session_state.get(PERFORMANCE_STATE_KEY, {})
        if metrics:
            performance_rows = [
                {
                    "Vorgang": label,
                    "Zeit (ms)": value.get("milliseconds", 0.0),
                    "Quelle": "Cache" if value.get("cached") else "Neu berechnet",
                    "Gemessen": value.get("measured_at", ""),
                }
                for label, value in metrics.items()
            ]
            st.dataframe(
                pd.DataFrame(performance_rows),
                hide_index=True,
                width="stretch",
            )
        else:
            st.info("Noch keine Messwerte vorhanden.")

        if st.button("Kontodaten jetzt neu laden", width="stretch"):
            invalidate_account_cache()
            st.rerun()
